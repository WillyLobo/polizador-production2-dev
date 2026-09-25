"""
Armado del .docx de una resolucion de certificado de obra.

Reusa las primitivas de formato de `secretariador/docx_builder.py` (fuente,
espaciados, sangrias de considerandos, negritas de los encabezados) en vez de
duplicarlas: ese modulo no importa nada de `secretariador` y los imports
cruzados entre apps ya existen en ambas direcciones. La diferencia con
viaticos es de donde sale el contenido -- alla el cuerpo tiene una estructura
fija conocida por el codigo, aca es una lista de bloques cargada por el
usuario, asi que el armado es un despacho por `clase`.

El documento base es el mismo membrete institucional que usa viaticos:
`secretariador.models.EncabezadoDocumento.vigente()`, un .docx de cuerpo vacio
que aporta header/footer/seccion/margenes.
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from secretariador.docx_builder import add_considerando, add_heading_bloque, add_heading_inline

# Ordinales de los articulos. La numeracion se deriva de la posicion del bloque
# dentro de la lista, no se guarda en la base: `secretariador/docx_builder.py`
# ya muestra lo que pasa cuando se escriben a mano (ahi conviven "º" y "°" en
# etiquetas contiguas), y ademas reordenar bloques renumeraria mal.
ORDINAL = "º"

CLASES = ("visto", "considerando", "resuelve", "articulo", "bloque")

# Encabezado fijo que separa los considerandos del articulado. No es editable
# porque no depende de los datos del certificado: es la formula del organismo.
ENCABEZADO_RESUELVE = "EL PRESIDENTE DEL INSTITUTO PROVINCIAL DE DESARROLLO URBANO Y VIVIENDA"


class ResolucionDocxError(Exception):
    """No se puede armar el documento (falta el membrete, por ejemplo)."""


def etiqueta_articulo(numero):
    return f"Artículo {numero}{ORDINAL}: "


def numerar_articulos(bloques):
    """Devuelve la lista de bloques con la etiqueta de cada articulo ya resuelta.

    Se respeta `label` si el bloque trae uno (las excepciones reales: "Artículo
    1º y 2º:", "De forma."); si no, se deriva del contador de articulos."""
    numerados = []
    contador = 0
    for bloque in bloques or []:
        bloque = dict(bloque)
        if bloque.get("clase") == "articulo":
            contador += 1
            if not bloque.get("label"):
                bloque["label"] = etiqueta_articulo(contador)
        numerados.append(bloque)
    return numerados


def build_resolucion_desde_bloques(base_docx_file, bloques):
    """Arma el cuerpo de la resolucion sobre `base_docx_file` (path o file-like:
    el .docx vigente de EncabezadoDocumento) y devuelve un BytesIO.

    `bloques` es la lista ya renderizada (ver `carga.resolucion_texto`), con el
    texto final -- aca no se evalua ninguna plantilla."""
    doc = Document(base_docx_file)

    for bloque in numerar_articulos(bloques):
        clase = bloque.get("clase", "considerando")
        texto = bloque.get("texto", "")
        label = bloque.get("label", "")

        if clase == "visto":
            add_heading_inline(doc, "VISTO: ", texto)
        elif clase == "considerando":
            add_considerando(doc, texto)
        elif clase == "resuelve":
            # Dos parrafos centrados, con keep_with_next para que Word no corte
            # la pagina entre el encabezado y el "RESUELVE:".
            add_heading_bloque(
                doc, ENCABEZADO_RESUELVE, jc=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True
            )
            add_heading_bloque(
                doc, "RESUELVE:", jc=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=None
            )
        elif clase == "articulo":
            add_heading_inline(doc, label, texto, before=240, after=0)
        elif clase == "bloque":
            add_heading_bloque(doc, texto)
        else:
            raise ResolucionDocxError(f"Clase de bloque desconocida: {clase!r}")

    salida = io.BytesIO()
    doc.save(salida)
    salida.seek(0)
    return salida


def build_resolucion_certificado(bloques):
    """Como `build_resolucion_desde_bloques`, pero resolviendo el membrete
    vigente. Lanza `ResolucionDocxError` con un mensaje accionable si todavia
    no se subio ninguno (mismo criterio que `revisar_texto_actuacion`)."""
    from secretariador.models import EncabezadoDocumento

    encabezado = EncabezadoDocumento.vigente()
    if encabezado is None:
        raise ResolucionDocxError(
            "No hay ningún documento base cargado. Subí uno desde "
            '"Actualizar Encabezado" antes de generar resoluciones.'
        )
    with encabezado.encabezadodocumento_archivo.open("rb") as f:
        base = io.BytesIO(f.read())
    return build_resolucion_desde_bloques(base, bloques)
