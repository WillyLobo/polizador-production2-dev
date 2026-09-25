import io
from datetime import date
from decimal import Decimal

from docx import Document

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db.utils import IntegrityError
from django.test import TestCase
from django.urls import reverse

from carga import certificacion, ley27397, resolucion_docx, resolucion_texto
from carga.certificacion import construir_certificados_desde_foja, generar_certificados_desde_foja
from carga.forms.fojademedicionforms import FojaDeMedicionForm
from carga.forms.contratotramopagoforms import ContratoTramoPagoFormset
from carga.models import (
    Certificado,
    CertificadoFinanciamiento,
    CertificadoRubro,
    Contrato,
    ContratoMonto,
    ContratoTramoPago,
    Empresa,
    FojaDeMedicion,
    FojaDeMedicionItem,
    Obra,
    PlanDeTrabajos,
    PlanDeTrabajosEtapa,
    PlanDeTrabajosEtapaItem,
    PlanDeTrabajosItem,
    PlanDeTrabajosRubro,
    Programa,
    TextoResolucionCertificado,
    Uvi,
)


class UviPesosEquivalentesTests(TestCase):
    def setUp(self):
        Uvi.objects.create(uvi_fecha=date(2026, 1, 1), uvi_valor=Decimal("100.00"))
        Uvi.objects.create(uvi_fecha=date(2026, 3, 1), uvi_valor=Decimal("120.00"))

    def test_usa_la_cotizacion_exacta_de_la_fecha(self):
        resultado = Uvi.pesos_equivalentes(Decimal("10"), date(2026, 3, 1))
        self.assertEqual(resultado, Decimal("1200.00"))

    def test_usa_la_cotizacion_anterior_mas_cercana(self):
        resultado = Uvi.pesos_equivalentes(Decimal("10"), date(2026, 2, 15))
        self.assertEqual(resultado, Decimal("1000.00"))

    def test_sin_cotizacion_previa_devuelve_none(self):
        resultado = Uvi.pesos_equivalentes(Decimal("10"), date(2025, 1, 1))
        self.assertIsNone(resultado)

    def test_sin_monto_o_fecha_devuelve_none(self):
        self.assertIsNone(Uvi.pesos_equivalentes(None, date(2026, 3, 1)))
        self.assertIsNone(Uvi.pesos_equivalentes(Decimal("10"), None))


class PlanDeTrabajosEtapaTestCase(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test",
            obra_empresa=empresa,
            obra_programa=programa,
            obra_expediente="EXP-1",
        )
        self.plan = PlanDeTrabajos.objects.create(
            trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1)
        )

    def _crear_rubro(self, **kwargs):
        defaults = {"rubro_plan": self.plan, "rubro_nombre": "Vivienda", "rubro_presupuesto": Decimal("1000")}
        defaults.update(kwargs)
        return PlanDeTrabajosRubro.objects.create(**defaults)

    def _crear_contratomonto(self, pesos=Decimal("0"), uvi=Decimal("0"), uvi_fecha=None):
        certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        financiamiento = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))
        return ContratoMonto.objects.create(
            contratomonto_contrato=contrato,
            contratomonto_rubro=certificado_rubro,
            contratomonto_financiamiento=financiamiento,
            contratomonto_pesos=pesos,
            contratomonto_uvi=uvi,
            contratomonto_uvi_fecha=uvi_fecha,
        )

    def test_monto_base_pesos_usa_rubro_presupuesto_sin_contratomonto(self):
        rubro = self._crear_rubro(rubro_presupuesto=Decimal("5000"))
        self.assertEqual(rubro.monto_base_pesos(), Decimal("5000"))

    def test_monto_base_pesos_usa_contratomonto_en_pesos(self):
        contratomonto = self._crear_contratomonto(pesos=Decimal("8000"))
        rubro = self._crear_rubro(rubro_contratomonto=contratomonto)
        self.assertEqual(rubro.monto_base_pesos(), Decimal("8000"))

    def test_monto_base_pesos_convierte_contratomonto_en_uvi(self):
        Uvi.objects.create(uvi_fecha=date(2026, 1, 1), uvi_valor=Decimal("100"))
        contratomonto = self._crear_contratomonto(uvi=Decimal("50"), uvi_fecha=date(2026, 1, 1))
        rubro = self._crear_rubro(rubro_contratomonto=contratomonto)
        self.assertEqual(rubro.monto_base_pesos(), Decimal("5000"))

    def test_etapa_numero_y_fecha_se_asignan_secuencialmente(self):
        rubro = self._crear_rubro()
        item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )

        etapa1 = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro)
        etapa2 = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro)

        self.assertEqual(etapa1.etapa_numero, 1)
        self.assertEqual(etapa1.etapa_fecha, date(2026, 1, 1))
        self.assertEqual(etapa2.etapa_numero, 2)
        self.assertEqual(etapa2.etapa_fecha, date(2026, 2, 1))

    def test_etapaitem_acumula_contra_la_etapa_anterior(self):
        rubro = self._crear_rubro()
        item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        etapa1 = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro)
        etapa2 = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro)

        etapaitem1 = PlanDeTrabajosEtapaItem.objects.create(
            etapaitem_etapa=etapa1, etapaitem_planitem=item, etapaitem_pct_proyectado_mes=Decimal("30")
        )
        etapaitem2 = PlanDeTrabajosEtapaItem.objects.create(
            etapaitem_etapa=etapa2, etapaitem_planitem=item, etapaitem_pct_proyectado_mes=Decimal("20")
        )

        self.assertEqual(etapaitem1.etapaitem_pct_proyectado_acumulado, Decimal("30"))
        self.assertEqual(etapaitem2.etapaitem_pct_proyectado_acumulado, Decimal("50"))

    def test_numeracion_continua_a_traves_de_una_reprogramacion(self):
        """Si el rubro se reprograma, las nuevas etapas continúan numerando (no resetean a 1)."""
        rubro_viejo = self._crear_rubro()
        PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro_viejo)
        PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro_viejo)

        plan_nuevo = PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 6, 1))
        rubro_nuevo = PlanDeTrabajosRubro.objects.create(
            rubro_plan=plan_nuevo,
            rubro_nombre="Vivienda",
            rubro_presupuesto=Decimal("1000"),
            rubro_anterior=rubro_viejo,
        )

        etapa3 = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro_nuevo)

        self.assertEqual(etapa3.etapa_numero, 3)
        # La fecha sigue siendo consecutiva a partir de la última etapa de la cadena,
        # no de trabajos_fecha del nuevo plan.
        self.assertEqual(etapa3.etapa_fecha, date(2026, 3, 1))

    def test_etapa_monto_pesos(self):
        contratomonto = self._crear_contratomonto(pesos=Decimal("10000"))
        rubro = self._crear_rubro(rubro_contratomonto=contratomonto)
        item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        etapa = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro)
        PlanDeTrabajosEtapaItem.objects.create(
            etapaitem_etapa=etapa, etapaitem_planitem=item, etapaitem_pct_proyectado_mes=Decimal("25")
        )

        self.assertEqual(etapa.etapa_monto_pesos(), Decimal("2500.00"))


class MatrizPlanDeTrabajosEtapaFormValidationTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-2"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=obra, trabajos_fecha=date(2026, 1, 1))
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan, rubro_nombre="Vivienda", rubro_presupuesto=Decimal("1000")
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("40")
        )

    def test_rechaza_si_la_suma_de_columnas_supera_la_incidencia_del_item(self):
        from carga.forms.plandetrabajosetapaforms import build_matriz_form

        form_class = build_matriz_form([self.item], total_columns=2, anterior_map={})
        form = form_class({f"item_{self.item.pk}_col_0": "30", f"item_{self.item.pk}_col_1": "20"})
        self.assertFalse(form.is_valid())

    def test_acepta_si_la_suma_de_columnas_no_supera_la_incidencia(self):
        from carga.forms.plandetrabajosetapaforms import build_matriz_form

        form_class = build_matriz_form([self.item], total_columns=2, anterior_map={})
        form = form_class({f"item_{self.item.pk}_col_0": "20", f"item_{self.item.pk}_col_1": "20"})
        self.assertTrue(form.is_valid())

    def test_rechaza_si_una_columna_supera_el_100_por_ciento_entre_items(self):
        from carga.forms.plandetrabajosetapaforms import build_matriz_form

        item2 = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 2", planitem_incidencia_pct=Decimal("100")
        )
        form_class = build_matriz_form([self.item, item2], total_columns=1, anterior_map={})
        form = form_class({f"item_{self.item.pk}_col_0": "40", f"item_{item2.pk}_col_0": "70"})
        self.assertFalse(form.is_valid())


class PlanDeTrabajosEtapaMatrizViewTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-3"
        )
        self.plan = PlanDeTrabajos.objects.create(
            trabajos_obra=obra, trabajos_fecha=date(2026, 1, 1), trabajos_meses=2
        )
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan, rubro_nombre="Vivienda", rubro_presupuesto=Decimal("1000")
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        user = get_user_model().objects.create_superuser(username="tester", password="testpass123")
        self.client.force_login(user)
        self.url = reverse("carga:plandetrabajosetapa-matriz", kwargs={"pk": self.rubro.pk})

    def _post_data(self, valores):
        """`valores`: un valor de % Proyectado del Mes por columna, para el único item."""
        return {f"item_{self.item.pk}_col_{col}": str(valor) for col, valor in enumerate(valores)}

    def test_get_muestra_tantas_columnas_como_trabajos_meses(self):
        response = self.client.get(self.url)
        self.assertContains(response, f'name="item_{self.item.pk}_col_0"')
        self.assertContains(response, f'name="item_{self.item.pk}_col_1"')
        self.assertNotContains(response, f'name="item_{self.item.pk}_col_2"')

    def test_post_crea_una_etapa_por_columna_con_numeracion_y_fecha_correctas(self):
        response = self.client.post(self.url, self._post_data(["30", "20"]))

        self.assertRedirects(
            response, reverse("carga:estado-obra", kwargs={"pk": self.plan.trabajos_obra_id})
        )
        etapas = list(PlanDeTrabajosEtapa.objects.order_by("etapa_numero"))
        self.assertEqual([e.etapa_numero for e in etapas], [1, 2])
        self.assertEqual(etapas[0].etapa_fecha, date(2026, 1, 1))
        self.assertEqual(etapas[1].etapa_fecha, date(2026, 2, 1))

        items = PlanDeTrabajosEtapaItem.objects.filter(
            etapaitem_etapa__in=etapas
        ).order_by("etapaitem_etapa__etapa_numero")
        self.assertEqual([i.etapaitem_pct_proyectado_mes for i in items], [Decimal("30"), Decimal("20")])
        self.assertEqual([i.etapaitem_pct_proyectado_acumulado for i in items], [Decimal("30"), Decimal("50")])

    def test_reabrir_y_modificar_una_columna_anterior_recalcula_el_acumulado_de_las_siguientes(self):
        self.client.post(self.url, self._post_data(["30", "20"]))

        response = self.client.post(self.url, self._post_data(["10", "20"]))

        self.assertEqual(response.status_code, 302)
        items = PlanDeTrabajosEtapaItem.objects.filter(
            etapaitem_planitem=self.item
        ).order_by("etapaitem_etapa__etapa_numero")
        self.assertEqual([i.etapaitem_pct_proyectado_mes for i in items], [Decimal("10"), Decimal("20")])
        self.assertEqual([i.etapaitem_pct_proyectado_acumulado for i in items], [Decimal("10"), Decimal("30")])

    def test_post_rechaza_si_la_suma_total_supera_la_incidencia_del_item(self):
        self.item.planitem_incidencia_pct = Decimal("40")
        self.item.save()

        response = self.client.post(self.url, self._post_data(["30", "20"]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(PlanDeTrabajosEtapa.objects.count(), 0)

    def test_completar_etapas_con_una_ya_guardada_no_duplica_el_acumulado_anterior(self):
        """Si ya existe 1 etapa guardada y se reabre la matriz para completar las
        restantes, la etapa existente no debe contarse dos veces (una via el
        acumulado anterior y otra como columna de la matriz)."""
        self.item.planitem_incidencia_pct = Decimal("40")
        self.item.save()
        self.plan.trabajos_meses = 1
        self.plan.save()
        self.client.post(self.url, self._post_data(["30"]))
        self.assertEqual(PlanDeTrabajosEtapa.objects.count(), 1)

        self.plan.trabajos_meses = 3
        self.plan.save()
        response = self.client.post(self.url, self._post_data(["30", "5", "5"]))

        self.assertRedirects(
            response, reverse("carga:estado-obra", kwargs={"pk": self.plan.trabajos_obra_id})
        )
        self.assertEqual(PlanDeTrabajosEtapa.objects.count(), 3)


class FojaDeMedicionNumeracionTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test",
            obra_empresa=empresa,
            obra_programa=programa,
            obra_expediente="EXP-1",
        )
        self.plan = PlanDeTrabajos.objects.create(
            trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1)
        )
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan,
            rubro_nombre="Vivienda",
            rubro_presupuesto=Decimal("1000"),
            rubro_foja_numero_inicial=5,
        )
        self.certificado_rubro_db = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )

    def _crear_certificado(self, **kwargs):
        defaults = {
            "certificado_obra": self.obra,
            "certificado_expediente": "EXP-CERT-1",
            "certificado_rubro_db": self.certificado_rubro_db,
        }
        defaults.update(kwargs)
        return Certificado.objects.create(**defaults)

    def test_primera_foja_real_toma_el_numero_inicial_del_rubro(self):
        foja = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=date(2026, 1, 1))
        self.assertEqual(foja.foja_numero, 5)

    def test_foja_legacy_con_numero_manual_y_vinculo_a_certificado(self):
        certificado = self._crear_certificado()

        foja = FojaDeMedicion(foja_rubro=self.rubro, foja_periodo=date(2020, 1, 1), foja_legacy=True)
        foja.foja_numero = 2
        foja.save()

        certificado.certificado_foja = foja
        certificado.save(update_fields=["certificado_foja"])

        foja.refresh_from_db()
        certificado.refresh_from_db()
        self.assertEqual(foja.foja_numero, 2)
        self.assertTrue(foja.foja_legacy)
        self.assertEqual(certificado.certificado_foja_id, foja.pk)

    def test_foja_real_despues_de_legacy_continua_la_secuencia(self):
        foja_legacy = FojaDeMedicion(foja_rubro=self.rubro, foja_periodo=date(2020, 1, 1), foja_legacy=True)
        foja_legacy.foja_numero = 3
        foja_legacy.save()

        foja_legacy_2 = FojaDeMedicion(foja_rubro=self.rubro, foja_periodo=date(2020, 2, 1), foja_legacy=True)
        foja_legacy_2.foja_numero = 4
        foja_legacy_2.save()

        foja_real = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=date(2026, 1, 1))
        self.assertEqual(foja_real.foja_numero, 5)

    def _form_data(self, foja_numero_manual):
        return {
            "foja_rubro": self.rubro.pk,
            "foja_legacy": True,
            "foja_numero_manual": foja_numero_manual,
            "foja_periodo": "2020-02-01",
            "foja_fecha": "2020-02-01",
        }

    def test_form_rechaza_numero_duplicado_en_modo_legacy(self):
        foja_existente = FojaDeMedicion(foja_rubro=self.rubro, foja_periodo=date(2020, 1, 1), foja_legacy=True)
        foja_existente.foja_numero = 2
        foja_existente.save()

        form = FojaDeMedicionForm(data=self._form_data(2))

        self.assertFalse(form.is_valid())
        self.assertIn("foja_numero_manual", form.errors)

    def test_form_rechaza_numero_fuera_de_rango_en_modo_legacy(self):
        form = FojaDeMedicionForm(data=self._form_data(self.rubro.rubro_foja_numero_inicial))

        self.assertFalse(form.is_valid())
        self.assertIn("foja_numero_manual", form.errors)

    def _form_data_no_legacy(self):
        return {
            "foja_rubro": self.rubro.pk,
            "foja_periodo": "2026-01-01",
            "foja_fecha": "2026-01-01",
        }

    def test_form_rechaza_foja_real_si_el_rubro_no_tiene_etapas(self):
        form = FojaDeMedicionForm(data=self._form_data_no_legacy())

        self.assertFalse(form.is_valid())
        self.assertIn("foja_rubro", form.errors)

    def test_form_acepta_foja_real_si_el_rubro_tiene_etapas(self):
        PlanDeTrabajosEtapa.objects.create(etapa_rubro=self.rubro)

        form = FojaDeMedicionForm(data=self._form_data_no_legacy())

        self.assertTrue(form.is_valid(), form.errors)

    def test_form_acepta_foja_legacy_aunque_el_rubro_no_tenga_etapas(self):
        form = FojaDeMedicionForm(data=self._form_data(2))

        self.assertNotIn("foja_rubro", form.errors)


class FojaDeMedicionAcumuladoCascadaTests(TestCase):
    """Bug real (Obra 1242): al corregir el % de una Foja ya cargada, las Fojas
    posteriores quedaban con `fojaitem_pct_acumulado` desactualizado (calculado una
    sola vez al guardarse, no derivado), pudiendo terminar mostrando un % acumulado
    menor al de la Foja anterior."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-CASCADA"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1))
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan, rubro_nombre="Vivienda", rubro_presupuesto=Decimal("1000")
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )

    def _crear_foja(self, periodo, pct_avance):
        foja = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=periodo)
        item = FojaDeMedicionItem.objects.create(
            fojaitem_foja=foja, fojaitem_planitem=self.item, fojaitem_pct_avance_mes=Decimal(pct_avance)
        )
        return foja, item

    def test_editar_una_foja_anterior_recalcula_en_cascada_la_siguiente(self):
        _, item10 = self._crear_foja(date(2026, 1, 1), "2.210")
        _, item11 = self._crear_foja(date(2026, 2, 1), "0.200")

        item11.refresh_from_db()
        self.assertEqual(item11.fojaitem_pct_acumulado, Decimal("2.410"))

        # Corrección posterior de Foja 10 (ej. se detectó que el avance real fue mayor).
        item10.fojaitem_pct_avance_mes = Decimal("2.780")
        item10.save()

        item11.refresh_from_db()
        self.assertEqual(item10.fojaitem_pct_acumulado, Decimal("2.780"))
        self.assertEqual(item11.fojaitem_pct_acumulado, Decimal("2.980"))
        self.assertGreaterEqual(item11.fojaitem_pct_acumulado, item10.fojaitem_pct_acumulado)

    def test_la_cascada_se_propaga_a_traves_de_varias_fojas(self):
        _, item1 = self._crear_foja(date(2026, 1, 1), "10")
        _, item2 = self._crear_foja(date(2026, 2, 1), "10")
        _, item3 = self._crear_foja(date(2026, 3, 1), "10")

        item1.fojaitem_pct_avance_mes = Decimal("20")
        item1.save()

        item2.refresh_from_db()
        item3.refresh_from_db()
        self.assertEqual(item2.fojaitem_pct_acumulado, Decimal("30"))
        self.assertEqual(item3.fojaitem_pct_acumulado, Decimal("40"))


class GenerarCertificadosDesdeFojaTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-1"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1))
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan,
            rubro_nombre="Vivienda",
            rubro_presupuesto=Decimal("1000"),
            rubro_certificado_rubro=self.certificado_rubro,
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        self.contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))

    def _crear_financiamiento(self, nombre="Nación", corto="N"):
        return CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre=nombre, certificadofinanciamiento_nombre_corto=corto
        )

    def _crear_contratomonto(self, financiamiento, pesos=Decimal("0"), uvi=Decimal("0")):
        return ContratoMonto.objects.create(
            contratomonto_contrato=self.contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=financiamiento,
            contratomonto_pesos=pesos,
            contratomonto_uvi=uvi,
        )

    def _crear_foja(self, periodo, pct_avance):
        foja = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=periodo)
        FojaDeMedicionItem.objects.create(
            fojaitem_foja=foja, fojaitem_planitem=self.item, fojaitem_pct_avance_mes=Decimal(pct_avance)
        )
        return foja

    def test_genera_un_certificado_por_cada_financiamiento(self):
        nacion = self._crear_financiamiento("Nación", "N")
        provincia = self._crear_financiamiento("Provincia", "P")
        self._crear_contratomonto(nacion, pesos=Decimal("100000"))
        self._crear_contratomonto(provincia, pesos=Decimal("50000"))
        foja = self._crear_foja(date(2026, 1, 1), "10")

        certificados = generar_certificados_desde_foja(foja, "EXP-CERT-1", date(2026, 2, 1))

        self.assertEqual(len(certificados), 2)
        montos = {c.certificado_financiamiento: c.certificado_monto_pesos for c in certificados}
        self.assertEqual(montos, {"N": Decimal("10000.00"), "P": Decimal("5000.00")})

    def test_falla_si_el_rubro_no_esta_vinculado_a_un_rubro_de_certificado(self):
        self.rubro.rubro_certificado_rubro = None
        self.rubro.save()
        foja = self._crear_foja(date(2026, 1, 1), "10")

        with self.assertRaises(ValidationError):
            construir_certificados_desde_foja(foja, "EXP", date(2026, 2, 1))

    def test_usa_rubro_contratomonto_directo_si_no_hay_rubro_certificado_rubro(self):
        """Caso más común (y el de los rubros cargados antes de esta feature): el rubro
        tiene un único ContratoMonto vinculado directamente vía rubro_contratomonto, sin
        necesidad de vincular también un Rubro de Certificado."""
        nacion = self._crear_financiamiento()
        contratomonto = self._crear_contratomonto(nacion, pesos=Decimal("100000"))
        self.rubro.rubro_certificado_rubro = None
        self.rubro.rubro_contratomonto = contratomonto
        self.rubro.save()
        foja = self._crear_foja(date(2026, 1, 1), "10")

        certificados = generar_certificados_desde_foja(foja, "EXP", date(2026, 2, 1))

        self.assertEqual(len(certificados), 1)
        self.assertEqual(certificados[0].certificado_financiamiento, "N")
        self.assertEqual(certificados[0].certificado_monto_pesos, Decimal("10000.00"))

    def test_falla_si_la_obra_no_tiene_contrato(self):
        self.contrato.delete()
        foja = self._crear_foja(date(2026, 1, 1), "10")

        with self.assertRaises(ValidationError):
            construir_certificados_desde_foja(foja, "EXP", date(2026, 2, 1))

    def test_solo_el_plan_vigente_puede_generar_certificados(self):
        nacion = self._crear_financiamiento()
        self._crear_contratomonto(nacion, pesos=Decimal("1000"))
        foja_vieja = self._crear_foja(date(2026, 1, 1), "10")

        # Reprograma el plan -> el plan viejo deja de ser vigente.
        PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 6, 1))

        with self.assertRaises(ValidationError):
            construir_certificados_desde_foja(foja_vieja, "EXP", date(2026, 2, 1))

    def test_copia_mes_pct_ante_pct_y_acum_pct_desde_la_foja(self):
        nacion = self._crear_financiamiento()
        self._crear_contratomonto(nacion, pesos=Decimal("1000"))
        foja1 = self._crear_foja(date(2026, 1, 1), "10")
        generar_certificados_desde_foja(foja1, "EXP-1", date(2026, 2, 1))
        foja2 = self._crear_foja(date(2026, 2, 1), "15")

        certificado = generar_certificados_desde_foja(foja2, "EXP-2", date(2026, 3, 1))[0]

        self.assertEqual(certificado.certificado_mes_pct, Decimal("15"))
        self.assertEqual(certificado.certificado_acum_pct, Decimal("25"))
        self.assertEqual(certificado.certificado_ante_pct, Decimal("10"))

    def test_certificado_tipo_y_foja_quedan_seteados(self):
        nacion = self._crear_financiamiento()
        self._crear_contratomonto(nacion, pesos=Decimal("1000"))
        foja = self._crear_foja(date(2026, 1, 1), "10")

        certificado = generar_certificados_desde_foja(foja, "EXP", date(2026, 2, 1))[0]

        self.assertEqual(certificado.certificado_tipo, "PARCIAL")
        self.assertEqual(certificado.certificado_foja_id, foja.pk)


class ResumenCertificacionMensualTests(TestCase):
    """certificacion.resumen_certificacion_mensual: cortes mes/anterior/total y
    fondo de reparo, sobre certificados creados directamente (sin pasar por
    generar_certificados_desde_foja, para aislar la lógica de agregación)."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test",
            obra_empresa=empresa,
            obra_programa=programa,
            obra_expediente="EXP-1",
            obra_contrato_nacion_pesos=Decimal("100000"),
        )
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        financiamiento_nacion = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))
        ContratoMonto.objects.create(
            contratomonto_contrato=contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=financiamiento_nacion,
            contratomonto_pesos=Decimal("100000"),
        )

    def _crear_certificado(self, tipo, fecha, monto_pesos, **kwargs):
        defaults = {
            "certificado_obra": self.obra,
            "certificado_tipo": tipo,
            "certificado_financiamiento": "N",
            "certificado_rubro_db": self.certificado_rubro,
            "certificado_expediente": "EXP-CERT",
            "certificado_fecha": fecha,
            "certificado_monto_pesos": monto_pesos,
        }
        defaults.update(kwargs)
        return Certificado.objects.create(**defaults)

    def test_corta_mes_anterior_y_total_para_certificado_de_obra(self):
        cert1 = self._crear_certificado("PARCIAL", date(2026, 1, 1), Decimal("10000"))
        cert2 = self._crear_certificado("PARCIAL", date(2026, 2, 1), Decimal("15000"))

        resumen = certificacion.resumen_certificacion_mensual(cert2)

        fila = resumen["certificado_obra"]
        self.assertEqual(fila["anterior_pesos"], Decimal("10000.00"))
        self.assertEqual(fila["mes_pesos"], Decimal("15000.00"))
        self.assertEqual(fila["total_pesos"], Decimal("25000.00"))
        self.assertEqual(fila["total_pct"], Decimal("25.000"))

    def test_anticipo_y_devolucion_no_se_mezclan_con_certificado_de_obra(self):
        self._crear_certificado("ANTICIPO", date(2026, 1, 15), Decimal("5000"))
        cert2 = self._crear_certificado("PARCIAL", date(2026, 2, 1), Decimal("15000"))

        resumen = certificacion.resumen_certificacion_mensual(cert2)

        self.assertEqual(resumen["anticipo"]["total_pesos"], Decimal("5000.00"))
        self.assertEqual(resumen["anticipo"]["mes_pesos"], Decimal("0.00"))
        self.assertEqual(resumen["certificado_obra"]["total_pesos"], Decimal("15000.00"))

    def test_fondo_de_reparo_5pct_sobre_bruto_sin_descontar_anticipo(self):
        cert1 = self._crear_certificado(
            "PARCIAL", date(2026, 1, 1), Decimal("10000"), certificado_descuento_anticipo_pesos=Decimal("2000")
        )
        cert2 = self._crear_certificado("PARCIAL", date(2026, 2, 1), Decimal("15000"))

        resumen = certificacion.resumen_certificacion_mensual(cert2)

        # 5% de 10000 (bruto, sin descontar los 2000 de anticipo) + 5% de 15000
        self.assertEqual(resumen["fondo_reparo"]["anterior_pesos"], Decimal("500.00"))
        self.assertEqual(resumen["fondo_reparo"]["mes_pesos"], Decimal("750.00"))
        self.assertEqual(resumen["fondo_reparo"]["total_pesos"], Decimal("1250.00"))
        self.assertEqual(
            resumen["total_general"]["total_pesos"],
            resumen["subtotal1"]["total_pesos"] - Decimal("1250.00"),
        )

    def test_fondo_de_reparo_no_aplica_a_anticipo(self):
        cert = self._crear_certificado("ANTICIPO", date(2026, 1, 1), Decimal("5000"))
        # certificado_fondoreparo_monto_*() ignoran el % cargado si el tipo es ANTICIPO.
        self.assertEqual(cert.certificado_fondoreparo_monto_pesos(), Decimal("0"))
        self.assertEqual(cert.certificado_fondoreparo_monto_uvi(), Decimal("0"))

        # Certificado.clean() además fuerza el % a 0 para este tipo (se ejerce al
        # pasar por full_clean(), como hacen las vistas de creación).
        cert.full_clean()
        self.assertEqual(cert.certificado_fondoreparo_pct, Decimal("0"))

    def test_anticipo_pct_se_expresa_contra_el_pool_del_financiamiento_no_contra_el_rubro(self):
        # Un segundo rubro bajo el mismo financiamiento agranda el pool de la obra
        # (150000) por encima del contrato del rubro puntual del certificado (100000,
        # Vivienda, seteado en setUp) — reproduce el caso real (Certificado 13974).
        otro_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Reconocimiento de Trabajos", certificadorubro_nombre_corto="R"
        )
        contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))
        ContratoMonto.objects.create(
            contratomonto_contrato=contrato,
            contratomonto_rubro=otro_rubro,
            contratomonto_financiamiento=CertificadoFinanciamiento.objects.get(certificadofinanciamiento_nombre_corto="N"),
            contratomonto_pesos=Decimal("50000"),
        )
        self.obra.refresh_from_db()
        self.assertEqual(self.obra.obra_contrato_nacion_pesos, Decimal("150000"))

        # 20% del pool (150000) = 30000, no 20% del rubro puntual (100000).
        cert = self._crear_certificado(
            "ANTICIPO", date(2026, 1, 1), Decimal("30000"), certificado_anticipo_pct=Decimal("20")
        )

        resumen = certificacion.resumen_certificacion_mensual(cert)

        self.assertEqual(resumen["anticipo"]["total_pct"], Decimal("20.000"))


class Ley27397TestsBase(TestCase):
    """Fixtures compartidas: Obra con Plan de Trabajos de 3 Etapas proyectando 10%/mes
    cada una (Ene/Feb/Mar 2026), un único Item, y un ContratoMonto con componente UVI."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-L1"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1))
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.financiamiento = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan,
            rubro_nombre="Vivienda",
            rubro_presupuesto=Decimal("1000"),
            rubro_certificado_rubro=self.certificado_rubro,
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        self.contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2025, 12, 1))
        self.contratomonto = ContratoMonto.objects.create(
            contratomonto_contrato=self.contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_pesos=Decimal("0"),
            contratomonto_uvi=Decimal("10000"),
            contratomonto_uvi_fecha=date(2025, 12, 15),
        )
        # Cotización pactada del contrato + cotizaciones de fin de mes de Ene/Feb/Mar 2026.
        Uvi.objects.create(uvi_fecha=date(2025, 12, 15), uvi_valor=Decimal("100"))
        Uvi.objects.create(uvi_fecha=date(2026, 1, 31), uvi_valor=Decimal("110"))
        Uvi.objects.create(uvi_fecha=date(2026, 2, 28), uvi_valor=Decimal("120"))
        Uvi.objects.create(uvi_fecha=date(2026, 3, 31), uvi_valor=Decimal("130"))

        for _ in range(3):
            PlanDeTrabajosEtapa.objects.create(etapa_rubro=self.rubro)
        self.etapas = list(PlanDeTrabajosEtapa.objects.order_by("etapa_numero"))
        for etapa in self.etapas:
            PlanDeTrabajosEtapaItem.objects.create(
                etapaitem_etapa=etapa, etapaitem_planitem=self.item, etapaitem_pct_proyectado_mes=Decimal("10")
            )

    def _crear_foja(self, periodo, pct_avance):
        foja = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=periodo)
        FojaDeMedicionItem.objects.create(
            fojaitem_foja=foja, fojaitem_planitem=self.item, fojaitem_pct_avance_mes=Decimal(pct_avance)
        )
        return foja


class Ley27397Tests(Ley27397TestsBase):
    def test_mes_al_dia_usa_cotizacion_de_fin_de_su_propio_mes(self):
        foja = self._crear_foja(date(2026, 1, 1), "10")

        tramos = ley27397.resolver_tasas_periodo(foja, self.financiamiento)

        self.assertEqual(len(tramos), 1)
        self.assertEqual(tramos[0].pct, Decimal("10"))
        self.assertEqual(tramos[0].tasa_valor, Decimal("110"))
        self.assertEqual(tramos[0].tasa_fecha, date(2026, 1, 31))


    def test_primer_mes_en_atraso_usa_tasa_pactada_del_contrato(self):
        foja = self._crear_foja(date(2026, 1, 1), "5")

        tramos = ley27397.resolver_tasas_periodo(foja, self.financiamiento)

        self.assertEqual(len(tramos), 1)
        self.assertEqual(tramos[0].pct, Decimal("5"))
        self.assertEqual(tramos[0].tasa_valor, Decimal("100"))
        self.assertEqual(tramos[0].tasa_fecha, date(2025, 12, 15))

    def test_atraso_multi_mes_recuperado_en_una_sola_foja_usa_cotizacion_propia_de_cada_lote(self):
        """foja1 sólo abona parcialmente el lote de Enero (2 de 10, no llega al 90%,
        usa la tasa pactada de bootstrap). foja2 completa Enero y además liquida Febrero
        y Marzo completos en un solo envío: cada lote debe usar la cotización de fin de
        SU PROPIO mes al llegar a su 90%, no la heredada del bootstrap de foja1 (la
        corrección del algoritmo validada con el usuario)."""
        self._crear_foja(date(2026, 1, 1), "2")
        foja2 = self._crear_foja(date(2026, 2, 1), "28")

        tramos = ley27397.resolver_tasas_periodo(foja2, self.financiamiento)

        self.assertEqual(len(tramos), 3)
        self.assertEqual([t.pct for t in tramos], [Decimal("8"), Decimal("10"), Decimal("10")])
        # Ninguno de los tres tramos de foja2 debe quedar en 100 (la tasa heredada del
        # bootstrap de foja1): cada lote alcanza su propio 90% dentro de foja2 y usa su
        # propia cotización de fin de mes.
        self.assertEqual([t.tasa_valor for t in tramos], [Decimal("110"), Decimal("120"), Decimal("130")])
        self.assertEqual([t.tasa_fecha for t in tramos], [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)])

    def test_sin_proyeccion_lanza_error(self):
        rubro_sin_etapas = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan, rubro_nombre="Sin Etapas", rubro_certificado_rubro=self.certificado_rubro
        )
        item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=rubro_sin_etapas, planitem_nombre="Item", planitem_incidencia_pct=Decimal("100")
        )
        foja = FojaDeMedicion.objects.create(foja_rubro=rubro_sin_etapas, foja_periodo=date(2026, 1, 1))
        FojaDeMedicionItem.objects.create(fojaitem_foja=foja, fojaitem_planitem=item, fojaitem_pct_avance_mes=Decimal("5"))

        with self.assertRaises(ley27397.SinProyeccionError):
            ley27397.resolver_tasas_periodo(foja, self.financiamiento)

    def test_avance_supera_lo_proyectado_lanza_error(self):
        foja = self._crear_foja(date(2026, 1, 1), "35")

        with self.assertRaises(ley27397.ProyeccionInsuficienteError):
            ley27397.resolver_tasas_periodo(foja, self.financiamiento)

    def test_falta_cotizacion_de_fin_de_mes_lanza_error(self):
        Uvi.objects.filter(uvi_fecha=date(2026, 1, 31)).delete()
        foja = self._crear_foja(date(2026, 1, 1), "10")

        with self.assertRaises(ley27397.CotizacionFaltanteError):
            ley27397.resolver_tasas_periodo(foja, self.financiamiento)

    def test_contrato_sin_fecha_pactada_lanza_error(self):
        self.contratomonto.contratomonto_uvi_fecha = None
        self.contratomonto.save()
        foja = self._crear_foja(date(2026, 1, 1), "5")

        with self.assertRaises(ley27397.SinTasaContratoError):
            ley27397.resolver_tasas_periodo(foja, self.financiamiento)

    def test_sin_monto_base_uvi_si_la_etapa_reprogramada_no_esta_vinculada(self):
        """El lote pertenece a una versión anterior del rubro que nunca se vinculó a un
        Rubro de Certificado: resolver_tasas_periodo puede resolver la cotización (no
        necesita el ContratoMonto si el lote cumple su 90% de una), pero convertir ese
        tramo a pesos sí lo necesita."""
        rubro_viejo = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan, rubro_nombre="Vivienda (vieja)", rubro_certificado_rubro=None
        )
        item_viejo = PlanDeTrabajosItem.objects.create(
            planitem_rubro=rubro_viejo, planitem_nombre="Item", planitem_incidencia_pct=Decimal("100")
        )
        etapa_vieja = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro_viejo)
        PlanDeTrabajosEtapaItem.objects.create(
            etapaitem_etapa=etapa_vieja, etapaitem_planitem=item_viejo, etapaitem_pct_proyectado_mes=Decimal("10")
        )
        # Sin ninguna Foja aún sobre rubro_viejo: será la Foja nueva (de más abajo) la
        # primera en abonar (y completar) este lote, referenciando la etapa vieja
        # sin vínculo a Rubro de Certificado.

        rubro_nuevo = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan,
            rubro_nombre="Vivienda (reprogramada)",
            rubro_certificado_rubro=self.certificado_rubro,
            rubro_anterior=rubro_viejo,
        )
        item_nuevo = PlanDeTrabajosItem.objects.create(
            planitem_rubro=rubro_nuevo, planitem_nombre="Item", planitem_incidencia_pct=Decimal("100")
        )
        etapa_nueva = PlanDeTrabajosEtapa.objects.create(etapa_rubro=rubro_nuevo)
        PlanDeTrabajosEtapaItem.objects.create(
            etapaitem_etapa=etapa_nueva, etapaitem_planitem=item_nuevo, etapaitem_pct_proyectado_mes=Decimal("10")
        )
        # pct=10 completa exactamente (y sólo) el lote de la etapa vieja (target 10);
        # el lote de la etapa nueva queda sin tocar.
        foja_nueva = FojaDeMedicion.objects.create(foja_rubro=rubro_nuevo, foja_periodo=date(2026, 1, 1))
        FojaDeMedicionItem.objects.create(
            fojaitem_foja=foja_nueva, fojaitem_planitem=item_nuevo, fojaitem_pct_avance_mes=Decimal("10")
        )

        tramos = ley27397.resolver_tasas_periodo(foja_nueva, self.financiamiento)
        self.assertEqual(len(tramos), 1)
        self.assertEqual(tramos[0].lote, etapa_vieja)

        with self.assertRaises(ley27397.SinMontoBaseUviError):
            ley27397.tramos_a_pesos(tramos, self.financiamiento)


class ContratoMontoEntreObrasTests(TestCase):
    """CertificadoRubro/CertificadoFinanciamiento son catálogos COMPARTIDOS entre todas
    las obras (ej. "Vivienda"/"Nación" no son exclusivos de una obra en particular). Un
    bug real (obra 1653) mostró un monto ~3x mayor al esperado porque
    `_contratomonto_de_rubro` buscaba el ContratoMonto sólo por Rubro+Financiamiento, sin
    acotar por el Contrato vigente de la obra dueña del rubro — así que podía traer el
    ContratoMonto de una obra completamente distinta que compartiera el mismo Rubro de
    Certificado + Financiamiento (en ese caso, con un pk de Contrato menor, que "ganaba"
    bajo el ordering por default de ContratoMonto)."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.financiamiento = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )

        # Obra ajena, creada PRIMERO (pk de Contrato menor) con un ContratoMonto para el
        # mismo Rubro de Certificado + Financiamiento, pero un monto en UVI muy distinto.
        obra_ajena = Obra.objects.create(
            obra_nombre="Obra Ajena", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-AJENA"
        )
        contrato_ajeno = Contrato.objects.create(contrato_obra=obra_ajena, contrato_fecha=date(2020, 1, 1))
        ContratoMonto.objects.create(
            contratomonto_contrato=contrato_ajeno,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_uvi=Decimal("999999"),
        )

        # La obra bajo prueba, creada DESPUÉS (pk de Contrato mayor).
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-PROPIA"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1))
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan,
            rubro_nombre="Vivienda",
            rubro_presupuesto=Decimal("1000"),
            rubro_certificado_rubro=self.certificado_rubro,
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        contrato_propio = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))
        self.contratomonto = ContratoMonto.objects.create(
            contratomonto_contrato=contrato_propio,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_uvi=Decimal("10000"),
            contratomonto_uvi_fecha=date(2025, 12, 15),
        )
        Uvi.objects.create(uvi_fecha=date(2025, 12, 15), uvi_valor=Decimal("100"))
        Uvi.objects.create(uvi_fecha=date(2026, 1, 31), uvi_valor=Decimal("110"))
        etapa = PlanDeTrabajosEtapa.objects.create(etapa_rubro=self.rubro)
        PlanDeTrabajosEtapaItem.objects.create(
            etapaitem_etapa=etapa, etapaitem_planitem=self.item, etapaitem_pct_proyectado_mes=Decimal("10")
        )

    def test_no_confunde_contratomonto_de_otra_obra_con_el_mismo_rubro_certificado(self):
        foja = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=date(2026, 1, 1))
        FojaDeMedicionItem.objects.create(
            fojaitem_foja=foja, fojaitem_planitem=self.item, fojaitem_pct_avance_mes=Decimal("10")
        )

        tramos = ley27397.resolver_tasas_periodo(foja, self.financiamiento)
        total_pesos = ley27397.tramos_a_pesos(tramos, self.financiamiento)

        # 10% de 10000 UVI (el ContratoMonto de ESTA obra) = 1000 UVI * 110 = 110000, no
        # algo derivado de 999999 (el ContratoMonto de la obra ajena).
        self.assertEqual(total_pesos, Decimal("110000"))


class Ley27397IntegrationTests(Ley27397TestsBase):
    def test_financiamiento_sin_uvi_no_aplica_ley27397(self):
        self.contratomonto.contratomonto_uvi = Decimal("0")
        self.contratomonto.contratomonto_pesos = Decimal("100000")
        self.contratomonto.save()
        foja = self._crear_foja(date(2026, 1, 1), "10")

        certificado = generar_certificados_desde_foja(foja, "EXP", date(2026, 2, 1))[0]

        self.assertEqual(certificado.certificado_monto_pesos, Decimal("10000.00"))
        self.assertIsNone(certificado.certificado_ley27397_detalle)

    def test_certificados_ya_generados_no_cambian_de_valor(self):
        foja1 = self._crear_foja(date(2026, 1, 1), "10")
        certificado1 = generar_certificados_desde_foja(foja1, "EXP-1", date(2026, 2, 1))[0]
        monto_original = certificado1.certificado_monto_pesos

        foja2 = self._crear_foja(date(2026, 2, 1), "10")
        generar_certificados_desde_foja(foja2, "EXP-2", date(2026, 3, 1))

        certificado1.refresh_from_db()
        self.assertEqual(certificado1.certificado_monto_pesos, monto_original)


class AnticipoTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-ANT"
        )
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.financiamiento_uvi = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        self.financiamiento_pesos = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Provincia", certificadofinanciamiento_nombre_corto="P"
        )
        contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))
        ContratoMonto.objects.create(
            contratomonto_contrato=contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento_uvi,
            contratomonto_pesos=Decimal("0"),
            contratomonto_uvi=Decimal("10000"),
            contratomonto_uvi_fecha=date(2025, 12, 15),
        )
        ContratoMonto.objects.create(
            contratomonto_contrato=contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento_pesos,
            contratomonto_pesos=Decimal("200000"),
            contratomonto_uvi=Decimal("0"),
        )
        self.obra.refresh_from_db()
        Uvi.objects.create(uvi_fecha=date(2025, 12, 15), uvi_valor=Decimal("100"))
        Uvi.objects.create(uvi_fecha=date(2026, 2, 1), uvi_valor=Decimal("150"))

    def _crear_certificado(self, **kwargs):
        defaults = {
            "certificado_obra": self.obra,
            "certificado_financiamiento": "N",
            "certificado_rubro_db": self.certificado_rubro,
            "certificado_expediente": "EXP",
            "certificado_fecha": date(2026, 2, 1),
        }
        defaults.update(kwargs)
        return Certificado(**defaults)

    def test_calcular_monto_anticipo_con_uvi_usa_cotizacion_del_dia_de_creacion(self):
        certificado = self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_anticipo_pct=Decimal("10"), certificado_fecha=date(2026, 2, 1)
        )

        certificacion.calcular_monto_anticipo(certificado)

        self.assertEqual(certificado.certificado_monto_uvi, Decimal("1000.00"))
        self.assertEqual(certificado.certificado_monto_pesos, Decimal("150000.00"))

    def test_calcular_monto_anticipo_sin_uvi_calcula_pesos_directo(self):
        certificado = self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_financiamiento="P", certificado_anticipo_pct=Decimal("25")
        )

        certificacion.calcular_monto_anticipo(certificado)

        self.assertEqual(certificado.certificado_monto_uvi, Decimal("0"))
        self.assertEqual(certificado.certificado_monto_pesos, Decimal("50000.00"))

    def test_calcular_monto_anticipo_calcula_anterior_y_acumulado(self):
        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_anticipo_pct=Decimal("10")
        ).save()

        certificado = self._crear_certificado(certificado_tipo="ANTICIPO", certificado_anticipo_pct=Decimal("15"))
        certificacion.calcular_monto_anticipo(certificado)

        self.assertEqual(certificado.certificado_anticipo_anterior, Decimal("10.000"))
        self.assertEqual(certificado.certificado_anticipo_acumulado, Decimal("25.000"))

    def test_calcular_monto_anticipo_anterior_no_cuenta_otro_financiamiento(self):
        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_financiamiento="P", certificado_anticipo_pct=Decimal("10")
        ).save()

        certificado = self._crear_certificado(certificado_tipo="ANTICIPO", certificado_anticipo_pct=Decimal("15"))
        certificacion.calcular_monto_anticipo(certificado)

        self.assertEqual(certificado.certificado_anticipo_anterior, Decimal("0.000"))
        self.assertEqual(certificado.certificado_anticipo_acumulado, Decimal("15.000"))

    def test_saldo_pendiente_anticipo_baja_con_devoluciones_y_llega_a_cero(self):
        anticipo = self._crear_certificado(certificado_tipo="ANTICIPO", certificado_anticipo_pct=Decimal("20"))
        certificacion.calcular_monto_anticipo(anticipo)
        self.assertEqual(anticipo.certificado_anticipo_saldo_pct, Decimal("20.000"))
        anticipo.save()

        cert1 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("5000"))
        certificacion.aplicar_descuento_anticipo(cert1)
        # El acumulado (bruto otorgado) no baja...
        self.assertEqual(cert1.certificado_descuento_anticipo_pct, Decimal("20.000"))
        # ...pero el saldo pendiente (neto) sí, a medida que se recupera el anticipo.
        self.assertEqual(cert1.certificado_anticipo_saldo_pct, Decimal("10.000"))
        cert1.save()

        cert2 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("5000"))
        certificacion.aplicar_descuento_anticipo(cert2)

        self.assertEqual(cert2.certificado_anticipo_saldo_pct, Decimal("0.000"))
        # El acumulado bruto de anticipo, en cambio, se mantiene en el 20% otorgado.
        self.assertEqual(anticipo.certificado_anticipo_acumulado, Decimal("20.000"))

    def test_sin_anticipo_no_hay_descuento(self):
        certificado = self._crear_certificado(
            certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("1000"), certificado_monto_pesos=Decimal("100000")
        )

        certificacion.aplicar_descuento_anticipo(certificado)

        self.assertEqual(certificado.certificado_descuento_anticipo_uvi, Decimal("0.00"))
        self.assertEqual(certificado.certificado_descuento_anticipo_pesos, Decimal("0.00"))
        # Sin descuento no corresponde numerar como "Devolución de Anticipo".
        self.assertEqual(certificado.certificado_rubro_devanticipo, 0)

    def test_con_descuento_se_numera_como_devolucion_de_anticipo(self):
        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_monto_uvi=Decimal("1000"), certificado_monto_pesos=Decimal("100000")
        ).save()

        cert1 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("2000"))
        certificacion.aplicar_descuento_anticipo(cert1)
        self.assertEqual(cert1.certificado_rubro_devanticipo, 1)
        cert1.save()

        cert2 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("2000"))
        certificacion.aplicar_descuento_anticipo(cert2)
        self.assertEqual(cert2.certificado_rubro_devanticipo, 2)

    def test_tasa_de_descuento_es_estable_certificado_a_certificado_sin_anticipo_nuevo(self):
        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_monto_uvi=Decimal("1000"), certificado_monto_pesos=Decimal("100000")
        ).save()

        cert1 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("3000"))
        certificacion.aplicar_descuento_anticipo(cert1)
        cert1.save()
        tasa1 = cert1.certificado_descuento_anticipo_uvi / cert1.certificado_monto_uvi

        cert2 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("2000"))
        certificacion.aplicar_descuento_anticipo(cert2)
        tasa2 = cert2.certificado_descuento_anticipo_uvi / cert2.certificado_monto_uvi

        self.assertEqual(tasa1, tasa2)

    def test_anticipo_nuevo_a_mitad_de_obra_sube_la_tasa_siguiente(self):
        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_monto_uvi=Decimal("1000"), certificado_monto_pesos=Decimal("100000")
        ).save()

        cert1 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("2000"))
        certificacion.aplicar_descuento_anticipo(cert1)
        cert1.save()
        tasa1 = cert1.certificado_descuento_anticipo_uvi / cert1.certificado_monto_uvi

        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_monto_uvi=Decimal("500"), certificado_monto_pesos=Decimal("50000")
        ).save()

        cert2 = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("2000"))
        certificacion.aplicar_descuento_anticipo(cert2)
        tasa2 = cert2.certificado_descuento_anticipo_uvi / cert2.certificado_monto_uvi

        self.assertGreater(tasa2, tasa1)

    def test_tasa_descuento_llega_a_uno_cuando_se_agota_el_saldo_a_certificar(self):
        self.assertEqual(certificacion._tasa_descuento(Decimal("0"), Decimal("100")), Decimal("0"))
        self.assertEqual(certificacion._tasa_descuento(Decimal("50"), Decimal("0")), Decimal("1"))
        self.assertEqual(certificacion._tasa_descuento(Decimal("50"), Decimal("100")), Decimal("0.5"))

    def test_anticipos_y_devoluciones_legacy_se_integran_al_calculo(self):
        self._crear_certificado(
            certificado_rubro_anticipo=Decimal("1"), certificado_monto_uvi=Decimal("500")
        ).save()
        self._crear_certificado(
            certificado_rubro_devanticipo=Decimal("1"), certificado_devolucion_monto_uvi=Decimal("200")
        ).save()

        cert = self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("1000"))
        certificacion.aplicar_descuento_anticipo(cert)

        # saldo pendiente = 500 (legacy otorgado) - 200 (legacy devuelto) = 300, sobre un
        # saldo a certificar de 10000 -> tasa 0.03 -> descuento = 1000 * 0.03 = 30.
        self.assertEqual(cert.certificado_descuento_anticipo_uvi, Decimal("30.00"))
        self.assertEqual(cert.certificado_descuento_anticipo_pct, Decimal("3.000"))

    def test_rechaza_anticipo_que_supere_el_30_por_ciento_pendiente(self):
        with self.assertRaises(ValidationError):
            certificacion.validar_anticipo_nuevo(self.obra, "N", Decimal("35"))

    def test_acepta_anticipo_si_lo_devuelto_libero_margen_bajo_el_30_por_ciento(self):
        self._crear_certificado(
            certificado_tipo="ANTICIPO", certificado_monto_uvi=Decimal("3000"), certificado_monto_pesos=Decimal("300000")
        ).save()
        self._crear_certificado(
            certificado_tipo="PARCIAL",
            certificado_monto_uvi=Decimal("2000"),
            certificado_descuento_anticipo_uvi=Decimal("1000"),
        ).save()

        # saldo pendiente = 3000 - 1000 = 2000 (20%); + 10% nuevo = 3000 (30%): no supera el tope.
        certificacion.validar_anticipo_nuevo(self.obra, "N", Decimal("10"))

    def test_rechaza_anticipo_sin_margen_en_saldo_a_certificar(self):
        self._crear_certificado(certificado_tipo="PARCIAL", certificado_monto_uvi=Decimal("9700")).save()

        # saldo a certificar = 10000 - 9700 = 300; un anticipo de 5% (500) no entra ahí,
        # aunque esté muy por debajo del tope del 30%.
        with self.assertRaises(ValidationError):
            certificacion.validar_anticipo_nuevo(self.obra, "N", Decimal("5"))


class HechoConsumadoTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-HC"
        )
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.financiamiento = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        self.contrato = Contrato.objects.create(contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1))
        self.contratomonto = ContratoMonto.objects.create(
            contratomonto_contrato=self.contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_pesos=Decimal("0"),
            contratomonto_uvi=Decimal("10000"),
            contratomonto_uvi_fecha=date(2025, 12, 15),
        )
        Uvi.objects.create(uvi_fecha=date(2025, 12, 15), uvi_valor=Decimal("100"))
        # Cotización de "hoy" muy distinta a la pactada, para poder distinguir cuál se usó.
        Uvi.objects.create(uvi_fecha=date(2026, 6, 1), uvi_valor=Decimal("500"))

    def _crear_certificado(self, **kwargs):
        defaults = {
            "certificado_obra": self.obra,
            "certificado_tipo": "HECHO_CONSUMADO",
            "certificado_contrato_origen": self.contrato,
            "certificado_financiamiento": "N",
            "certificado_rubro_db": self.certificado_rubro,
            "certificado_expediente": "EXP",
            "certificado_fecha": date(2026, 6, 1),
            "certificado_mes_pct": Decimal("10"),
        }
        defaults.update(kwargs)
        return Certificado(**defaults)

    def test_usa_la_cotizacion_de_la_fecha_del_certificado_no_la_pactada(self):
        certificado = self._crear_certificado()

        certificacion.calcular_monto_hecho_consumado(certificado)

        self.assertEqual(certificado.certificado_monto_uvi, Decimal("1000.00"))
        # 1000 UVI * 500 (cotización de certificado_fecha 2026-06-01), no * 100 (pactada 2025-12-15).
        self.assertEqual(certificado.certificado_monto_pesos, Decimal("500000.00"))

    def test_falla_sin_contratomonto_para_el_rubro_y_financiamiento(self):
        otro_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Infraestructura", certificadorubro_nombre_corto="I"
        )
        certificado = self._crear_certificado(certificado_rubro_db=otro_rubro)

        with self.assertRaises(ValidationError):
            certificacion.calcular_monto_hecho_consumado(certificado)

    def test_clean_exige_contrato_origen(self):
        certificado = self._crear_certificado(
            certificado_contrato_origen=None, certificado_monto_uvi=Decimal("1000"), certificado_monto_pesos=Decimal("100000")
        )

        with self.assertRaises(ValidationError):
            certificado.full_clean()

    def test_el_descuento_de_anticipo_tambien_aplica_a_hecho_consumado(self):
        Certificado.objects.create(
            certificado_obra=self.obra,
            certificado_tipo="ANTICIPO",
            certificado_financiamiento="N",
            certificado_rubro_db=self.certificado_rubro,
            certificado_expediente="EXP-ANT",
            certificado_monto_uvi=Decimal("1000"),
            certificado_monto_pesos=Decimal("100000"),
        )

        certificado = self._crear_certificado()
        certificacion.calcular_monto_hecho_consumado(certificado)
        certificacion.aplicar_descuento_anticipo(certificado)

        # saldo pendiente = 1000, saldo a certificar = 10000 -> tasa 0.1 -> descuento = 1000*0.1=100.
        self.assertEqual(certificado.certificado_descuento_anticipo_uvi, Decimal("100.00"))


class ContratoTramoPagoTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-TRAMO"
        )
        self.contrato = Contrato.objects.create(
            contrato_obra=obra, contrato_fecha=date(2026, 1, 1), contrato_certificacion_por_etapas=True
        )

    def test_numeracion_automatica_y_continua(self):
        tramo1 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("30"), tramo_pct_disparador=Decimal("0")
        )
        tramo2 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("70"), tramo_pct_disparador=Decimal("30")
        )
        self.assertEqual(tramo1.tramo_numero, 1)
        self.assertEqual(tramo2.tramo_numero, 2)

        tramo2.delete()
        tramo3 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("70"), tramo_pct_disparador=Decimal("30")
        )
        # Mismo criterio que auto_increment_foja_numero/etapa_numero: se basa en el último
        # tramo_numero existente, no en un contador global -> al borrar el 2 y crear uno
        # nuevo, vuelve a numerarse 2 (no 3).
        self.assertEqual(tramo3.tramo_numero, 2)

    def _formset(self, data):
        base_data = {
            "tramos_pago-TOTAL_FORMS": str(len(data)),
            "tramos_pago-INITIAL_FORMS": "0",
            "tramos_pago-MIN_NUM_FORMS": "0",
            "tramos_pago-MAX_NUM_FORMS": "1000",
        }
        for i, fila in enumerate(data):
            for campo, valor in fila.items():
                base_data[f"tramos_pago-{i}-{campo}"] = valor
        return ContratoTramoPagoFormset(base_data, instance=self.contrato)

    def test_formset_exige_que_la_suma_de_pct_pago_sea_100(self):
        formset = self._formset([
            {"tramo_pct_pago": "30", "tramo_pct_disparador": "0"},
            {"tramo_pct_pago": "30", "tramo_pct_disparador": "30"},
        ])
        self.assertFalse(formset.is_valid())

    def test_formset_acepta_suma_100(self):
        formset = self._formset([
            {"tramo_pct_pago": "30", "tramo_pct_disparador": "0"},
            {"tramo_pct_pago": "70", "tramo_pct_disparador": "30"},
        ])
        self.assertTrue(formset.is_valid())

    def test_formset_exige_disparador_no_decreciente(self):
        formset = self._formset([
            {"tramo_pct_pago": "70", "tramo_pct_disparador": "30"},
            {"tramo_pct_pago": "30", "tramo_pct_disparador": "0"},
        ])
        self.assertFalse(formset.is_valid())


class CertificadoEtapaTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-ETAPA"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=self.obra, trabajos_fecha=date(2026, 1, 1))
        self.certificado_rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan,
            rubro_nombre="Vivienda",
            rubro_presupuesto=Decimal("1000"),
            rubro_certificado_rubro=self.certificado_rubro,
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Item 1", planitem_incidencia_pct=Decimal("100")
        )
        self.contrato = Contrato.objects.create(
            contrato_obra=self.obra, contrato_fecha=date(2026, 1, 1), contrato_certificacion_por_etapas=True
        )
        self.financiamiento = CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        ContratoMonto.objects.create(
            contratomonto_contrato=self.contrato,
            contratomonto_rubro=self.certificado_rubro,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_pesos=Decimal("100000"),
        )
        self._crear_tramos()

    def _crear_tramos(self):
        self.tramo1 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("30"), tramo_pct_disparador=Decimal("0")
        )
        self.tramo2 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("30"), tramo_pct_disparador=Decimal("30")
        )
        self.tramo3 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("30"), tramo_pct_disparador=Decimal("60")
        )
        self.tramo4 = ContratoTramoPago.objects.create(
            tramo_contrato=self.contrato, tramo_pct_pago=Decimal("10"), tramo_pct_disparador=Decimal("90")
        )

    def _crear_foja(self, periodo, pct_avance):
        foja = FojaDeMedicion.objects.create(foja_rubro=self.rubro, foja_periodo=periodo)
        FojaDeMedicionItem.objects.create(
            fojaitem_foja=foja, fojaitem_planitem=self.item, fojaitem_pct_avance_mes=Decimal(pct_avance)
        )
        return foja

    def test_una_foja_que_cruza_exactamente_un_umbral_genera_un_certificado(self):
        foja = self._crear_foja(date(2026, 1, 1), "20")  # acumulado 20 -> sólo dispara tramo1 (umbral 0)

        certificados = generar_certificados_desde_foja(foja, "EXP-1", date(2026, 2, 1))

        self.assertEqual(len(certificados), 1)
        self.assertEqual(certificados[0].certificado_tipo, "ETAPA")
        self.assertEqual(certificados[0].certificado_contrato_tramo_id, self.tramo1.pk)
        self.assertEqual(certificados[0].certificado_monto_pesos, Decimal("30000.00"))

    def test_una_foja_que_salta_dos_umbrales_de_una_vez_genera_dos_certificados(self):
        foja1 = self._crear_foja(date(2026, 1, 1), "20")
        generar_certificados_desde_foja(foja1, "EXP-1", date(2026, 2, 1))

        foja2 = self._crear_foja(date(2026, 2, 1), "45")  # acumulado 65 -> cruza tramo2 (30) y tramo3 (60)
        certificados = generar_certificados_desde_foja(foja2, "EXP-2", date(2026, 3, 1))

        self.assertEqual(len(certificados), 2)
        tramos = {c.certificado_contrato_tramo_id for c in certificados}
        self.assertEqual(tramos, {self.tramo2.pk, self.tramo3.pk})

    def test_una_foja_que_no_cruza_ningun_umbral_nuevo_no_genera_nada(self):
        foja1 = self._crear_foja(date(2026, 1, 1), "20")
        generar_certificados_desde_foja(foja1, "EXP-1", date(2026, 2, 1))

        foja2 = self._crear_foja(date(2026, 2, 1), "5")  # acumulado 25, no llega al umbral 30 de tramo2
        certificados = construir_certificados_desde_foja(foja2, "EXP-2", date(2026, 3, 1))

        self.assertEqual(certificados, [])

    def test_un_tramo_ya_certificado_no_se_regenera(self):
        foja1 = self._crear_foja(date(2026, 1, 1), "50")  # acumulado 50 -> cruza tramo1 y tramo2
        generar_certificados_desde_foja(foja1, "EXP-1", date(2026, 2, 1))

        foja2 = self._crear_foja(date(2026, 2, 1), "5")  # acumulado 55, sigue superando tramo1/tramo2 pero ya están certificados
        certificados = construir_certificados_desde_foja(foja2, "EXP-2", date(2026, 3, 1))

        self.assertEqual(certificados, [])

    def test_contrato_por_etapas_nunca_genera_parcial(self):
        foja = self._crear_foja(date(2026, 1, 1), "20")

        certificados = construir_certificados_desde_foja(foja, "EXP", date(2026, 2, 1))

        self.assertTrue(all(c.certificado_tipo == "ETAPA" for c in certificados))

    def test_calcular_monto_etapa_suma_todos_los_rubros_del_contrato(self):
        otro_rubro_certificado = CertificadoRubro.objects.create(
            certificadorubro_nombre="Infraestructura", certificadorubro_nombre_corto="I"
        )
        ContratoMonto.objects.create(
            contratomonto_contrato=self.contrato,
            contratomonto_rubro=otro_rubro_certificado,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_pesos=Decimal("50000"),
        )
        certificado = Certificado(
            certificado_obra=self.obra,
            certificado_tipo="ETAPA",
            certificado_financiamiento="N",
            certificado_rubro_db=self.certificado_rubro,
            certificado_expediente="EXP",
            certificado_fecha=date(2026, 2, 1),
            certificado_contrato_tramo=self.tramo1,
        )

        certificacion.calcular_monto_etapa(certificado)

        # Base = 100000 (Vivienda) + 50000 (Infraestructura) = 150000; tramo1 = 30%.
        self.assertEqual(certificado.certificado_monto_pesos, Decimal("45000.00"))

    def test_clean_exige_foja_y_tramo_prohibe_contrato_origen(self):
        base = {
            "certificado_obra": self.obra,
            "certificado_tipo": "ETAPA",
            "certificado_financiamiento": "N",
            "certificado_rubro_db": self.certificado_rubro,
            "certificado_expediente": "EXP",
            "certificado_fecha": date(2026, 2, 1),
        }
        with self.assertRaises(ValidationError):
            Certificado(**base, certificado_contrato_tramo=self.tramo1).full_clean()

        foja = self._crear_foja(date(2026, 1, 1), "20")
        with self.assertRaises(ValidationError):
            Certificado(**base, certificado_foja=foja).full_clean()

        with self.assertRaises(ValidationError):
            Certificado(**base, certificado_foja=foja, certificado_contrato_tramo=self.tramo1, certificado_contrato_origen=self.contrato).full_clean()

    def test_obra_reconoce_certificado_etapa_como_avance(self):
        foja = self._crear_foja(date(2026, 1, 1), "20")
        generar_certificados_desde_foja(foja, "EXP-1", date(2026, 2, 1))

        self.obra.refresh_from_db()
        self.assertEqual(self.obra.obra_acum_pct(), Decimal("20.000"))

    def test_resumen_certificacion_usa_base_del_contrato_completo(self):
        otro_rubro_certificado = CertificadoRubro.objects.create(
            certificadorubro_nombre="Infraestructura", certificadorubro_nombre_corto="I"
        )
        ContratoMonto.objects.create(
            contratomonto_contrato=self.contrato,
            contratomonto_rubro=otro_rubro_certificado,
            contratomonto_financiamiento=self.financiamiento,
            contratomonto_pesos=Decimal("50000"),
        )
        foja = self._crear_foja(date(2026, 1, 1), "20")
        certificado = generar_certificados_desde_foja(foja, "EXP-1", date(2026, 2, 1))[0]

        resumen = certificacion.resumen_certificacion_mensual(certificado)

        # Base = 150000 (Vivienda + Infraestructura); monto del tramo = 45000 -> 30%,
        # no el % que daría contra un solo rubro (45000/100000 = 45%).
        self.assertEqual(resumen["etapa"]["total_pct"], Decimal("30.00"))
        # Fondo de Reparo (5% default) también debe reflejar el monto del certificado ETAPA.
        self.assertEqual(resumen["fondo_reparo"]["total"], Decimal("2250.00"))


class ResolucionTextoFiltrosTests(TestCase):
    """Los filtros se aplican sobre datos que pueden venir nulos (montos UVI de una
    obra financiada sólo en pesos, períodos sin cargar). Ninguno puede tirar una
    excepción: si lo hiciera, rompería la ficha del certificado."""

    def test_moneda_y_pesos_usan_separadores_es_ar(self):
        assert resolucion_texto.filtro_moneda(Decimal("1234567.8")) == "1.234.567,80"
        assert resolucion_texto.filtro_pesos(Decimal("1234567.8")) == "$1.234.567,80"

    def test_pct_agrega_el_simbolo(self):
        assert resolucion_texto.filtro_pct(Decimal("12.3456")) == "12,35 %"

    def test_letras_escribe_el_importe(self):
        assert resolucion_texto.filtro_letras(Decimal("1234.50")).startswith("Pesos mil doscientos treinta y cuatro")

    def test_fecha_en_formato_local(self):
        assert resolucion_texto.filtro_fecha(date(2026, 3, 9)) == "09/03/2026"

    def test_cuit_agrega_guiones(self):
        assert resolucion_texto.filtro_cuit("20123456789") == "20-12345678-9"

    def test_cuit_incompleto_se_devuelve_tal_cual(self):
        # A diferencia del filtro de numletras, que inventaría un "00-00000000-0":
        # en un texto que se firma es peor un CUIT falso que uno evidentemente incompleto.
        assert resolucion_texto.filtro_cuit("123") == "123"

    def test_ninguno_falla_con_none_ni_con_basura(self):
        for filtro in (resolucion_texto.filtro_moneda, resolucion_texto.filtro_pesos,
                       resolucion_texto.filtro_letras, resolucion_texto.filtro_fecha,
                       resolucion_texto.filtro_pct, resolucion_texto.filtro_cuit):
            assert filtro(None) == ""
            filtro("no es un número")

    def test_cero_no_se_confunde_con_vacio(self):
        assert resolucion_texto.filtro_pesos(Decimal("0")) == "$0,00"


class ResolucionTextoRenderTests(TestCase):
    def test_renderiza_con_filtros(self):
        texto, faltantes = resolucion_texto.render_bloque(
            "Por {{ c.monto|pesos }}.", {"c": {"monto": Decimal("1000")}}
        )
        assert texto == "Por $1.000,00."
        assert faltantes == set()

    def test_variable_inexistente_deja_marcador_y_se_reporta(self):
        texto, faltantes = resolucion_texto.render_bloque("Obra {{ obra.nombre }}", {})
        assert "«falta: obra»" in texto
        assert faltantes == {"obra"}

    def test_cadena_larga_de_atributos_no_explota(self):
        # ChainableUndefined: {{ obra.a.b }} tiene que degradar, no lanzar.
        texto, faltantes = resolucion_texto.render_bloque("{{ obra.a.b }}", {})
        assert "«falta:" in texto
        assert faltantes

    def test_error_de_sintaxis_es_error_de_dominio(self):
        with self.assertRaises(resolucion_texto.TextoResolucionError):
            resolucion_texto.render_bloque("{% if x %}sin cerrar", {})

    def test_compilar_bloque_detecta_sintaxis_invalida(self):
        with self.assertRaises(resolucion_texto.TextoResolucionError):
            resolucion_texto.compilar_bloque("{{ sin cerrar")
        # Un bloque válido no lanza.
        resolucion_texto.compilar_bloque("{{ certificado.expediente }}")

    def test_el_sandbox_bloquea_llamar_metodos_peligrosos(self):
        """La amenaza realista no es `__class__` (Jinja ya lo bloquea) sino invocar
        un método del objeto. El contexto son dicts proyectados, no modelos, así que
        `delete` ni siquiera existe, e invocarlo es un error de dominio -- no un
        marcador de "falta", porque no es un dato que se pueda completar."""
        with self.assertRaises(resolucion_texto.TextoResolucionError):
            resolucion_texto.render_bloque("{{ obra.delete() }}", {"obra": {"nombre": "X"}})

    def test_el_sandbox_bloquea_los_atributos_internos(self):
        texto, _ = resolucion_texto.render_bloque("{{ obra.__class__ }}", {"obra": {"nombre": "X"}})
        assert "«falta: __class__»" in texto

    def test_render_bloques_conserva_clase_y_label(self):
        bloques = [{"clase": "articulo", "label": "", "texto": "Monto {{ m|pesos }}"}]
        resueltos, _ = resolucion_texto.render_bloques(bloques, {"m": Decimal("5")})
        assert resueltos == [{"clase": "articulo", "label": "", "texto": "Monto $5,00"}]


class ResolucionDocxTests(TestCase):
    def _base_docx(self):
        """Documento base en memoria. No se usa `secretariador/media/*.docx` como
        fixture: esos traen cuerpo con tags de docxtpl, no son un encabezado de
        cuerpo vacío como el que sube EncabezadoDocumento."""
        base = io.BytesIO()
        Document().save(base)
        base.seek(0)
        return base

    def test_numera_los_articulos_por_posicion(self):
        bloques = [
            {"clase": "considerando", "texto": "Que..."},
            {"clase": "articulo", "texto": "Uno"},
            {"clase": "articulo", "texto": "Dos"},
        ]
        etiquetas = [b["label"] for b in resolucion_docx.numerar_articulos(bloques) if b["clase"] == "articulo"]
        assert etiquetas == ["Artículo 1º: ", "Artículo 2º: "]

    def test_respeta_la_etiqueta_manual(self):
        bloques = [
            {"clase": "articulo", "label": "", "texto": "Uno"},
            {"clase": "articulo", "label": "De forma: ", "texto": "Regístrese"},
        ]
        etiquetas = [b["label"] for b in resolucion_docx.numerar_articulos(bloques)]
        assert etiquetas == ["Artículo 1º: ", "De forma: "]

    def test_arma_el_documento_con_el_orden_y_las_negritas(self):
        bloques = [
            {"clase": "visto", "texto": "El Expediente N° 1."},
            {"clase": "bloque", "texto": "CONSIDERANDO:"},
            {"clase": "considerando", "texto": "Que corresponde;"},
            {"clase": "resuelve", "texto": ""},
            {"clase": "articulo", "texto": "Apruébase."},
        ]
        salida = resolucion_docx.build_resolucion_desde_bloques(self._base_docx(), bloques)
        parrafos = [p for p in Document(salida).paragraphs if p.text.strip()]
        textos = [p.text for p in parrafos]

        assert textos[0] == "VISTO: El Expediente N° 1."
        assert textos[1] == "CONSIDERANDO:"
        assert textos[2] == "Que corresponde;"
        assert "EL PRESIDENTE" in textos[3]
        assert textos[4] == "RESUELVE:"
        assert textos[5] == "Artículo 1º: Apruébase."
        # El considerando va sin negrita; los encabezados y la etiqueta del artículo sí.
        assert parrafos[2].runs[0].bold is False
        assert parrafos[1].runs[0].bold is True
        assert parrafos[5].runs[0].bold is True

    def test_clase_desconocida_es_error_explicito(self):
        with self.assertRaises(resolucion_docx.ResolucionDocxError):
            resolucion_docx.build_resolucion_desde_bloques(self._base_docx(), [{"clase": "inventada", "texto": "x"}])


class TextoResolucionAlcanceTests(TestCase):
    """Resolución del texto base que le toca a un certificado: gana el que coincide
    con el tipo exacto; si no hay, cae al comodín del programa y financiamiento."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        self.programa = Programa.objects.create(programa_nombre="Programa Test")
        self.otro_programa = Programa.objects.create(programa_nombre="Otro Programa")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=self.programa, obra_expediente="EXP-TR"
        )
        self.rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )

    def _texto(self, programa=None, financiamiento="N", tipo="", nombre="Texto"):
        return TextoResolucionCertificado.objects.create(
            textoresolucion_nombre=nombre,
            textoresolucion_programa=programa or self.programa,
            textoresolucion_financiamiento=financiamiento,
            textoresolucion_tipo=tipo,
            textoresolucion_bloques=[{"clase": "considerando", "label": "", "texto": nombre}],
        )

    def _certificado(self, tipo="PARCIAL", financiamiento="N"):
        return Certificado(
            certificado_obra=self.obra,
            certificado_tipo=tipo,
            certificado_financiamiento=financiamiento,
            certificado_rubro_db=self.rubro,
            certificado_expediente="EXP",
            certificado_fecha=date(2026, 6, 1),
        )

    def test_gana_el_tipo_exacto_sobre_el_comodin(self):
        self._texto(tipo="", nombre="Comodín")
        exacto = self._texto(tipo="PARCIAL", nombre="Parcial")

        assert TextoResolucionCertificado.para_certificado(self._certificado()) == exacto

    def test_cae_al_comodin_cuando_no_hay_texto_para_el_tipo(self):
        comodin = self._texto(tipo="", nombre="Comodín")
        self._texto(tipo="ANTICIPO", nombre="Anticipo")

        assert TextoResolucionCertificado.para_certificado(self._certificado(tipo="PARCIAL")) == comodin

    def test_sin_texto_para_el_alcance_devuelve_none(self):
        self._texto(financiamiento="P", tipo="")

        assert TextoResolucionCertificado.para_certificado(self._certificado(financiamiento="N")) is None

    def test_no_cruza_programas(self):
        self._texto(programa=self.otro_programa, tipo="")

        assert TextoResolucionCertificado.para_certificado(self._certificado()) is None

    def test_no_se_pueden_duplicar_dos_textos_para_el_mismo_alcance(self):
        self._texto(tipo="PARCIAL")
        with self.assertRaises(IntegrityError):
            self._texto(tipo="PARCIAL", nombre="Duplicado")

    def test_el_comodin_tambien_es_unico(self):
        # El centinela es "" y no NULL justamente para que la constraint lo dedupe:
        # con NULL, Postgres dejaría cargar comodines repetidos.
        self._texto(tipo="")
        with self.assertRaises(IntegrityError):
            self._texto(tipo="", nombre="Comodín duplicado")


class InvalidarTextoResolucionTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-INV"
        )
        rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.certificado = Certificado.objects.create(
            certificado_obra=obra,
            certificado_tipo="LEGACY",
            certificado_financiamiento="N",
            certificado_rubro_db=rubro,
            certificado_expediente="EXP-1",
            certificado_fecha=date(2026, 6, 1),
            certificado_monto_pesos=Decimal("1000"),
            certificado_texto_resolucion={"bloques": [{"clase": "considerando", "label": "", "texto": "Texto"}]},
        )

    def _recargar(self):
        return Certificado.objects.get(pk=self.certificado.pk)

    def test_cambiar_un_campo_que_el_texto_describe_lo_invalida(self):
        self.certificado.certificado_monto_pesos = Decimal("2000")
        self.certificado.save()

        assert self._recargar().certificado_texto_resolucion is None

    def test_cambiar_un_campo_ajeno_al_texto_no_lo_invalida(self):
        self.certificado.certificado_fecha_carga = date(2026, 7, 1)
        self.certificado.save()

        assert self._recargar().certificado_texto_resolucion is not None

    def test_la_instancia_en_memoria_queda_sincronizada(self):
        self.certificado.certificado_expediente = "EXP-2"
        self.certificado.save()

        assert self.certificado.certificado_texto_resolucion is None

    def test_no_agrega_una_fila_historica_extra(self):
        """La invalidación usa .update() justamente para no escribir en
        certificado_history: una edición tiene que dejar una sola fila, no dos."""
        antes = self.certificado.certificado_history.count()

        self.certificado.certificado_monto_pesos = Decimal("3000")
        self.certificado.save()

        assert self.certificado.certificado_history.count() == antes + 1

    def test_sin_snapshot_no_hace_nada(self):
        Certificado.objects.filter(pk=self.certificado.pk).update(certificado_texto_resolucion=None)
        certificado = self._recargar()
        certificado.certificado_monto_pesos = Decimal("4000")
        certificado.save()

        assert self._recargar().certificado_monto_pesos == Decimal("4000")


class TextoResolucionVistasTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Constructora Sur SRL", empresa_cuit="20123456789")
        self.programa = Programa.objects.create(programa_nombre="Techo Digno")
        self.obra = Obra.objects.create(
            obra_nombre="50 Viviendas en Resistencia", obra_empresa=empresa,
            obra_programa=self.programa, obra_expediente="EXP-OBRA",
        )
        self.rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        self.certificado = Certificado.objects.create(
            certificado_obra=self.obra,
            certificado_tipo="PARCIAL",
            certificado_financiamiento="N",
            certificado_rubro_db=self.rubro,
            certificado_expediente="EXP-CERT",
            certificado_periodo="Marzo 2026",
            certificado_fecha=date(2026, 3, 31),
            certificado_monto_pesos=Decimal("1500000"),
        )
        self.superuser = get_user_model().objects.create_superuser(username="admin", password="testpass123")
        self.sin_permisos = get_user_model().objects.create_user(username="pepe", password="testpass123")

    def _texto(self, tipo="PARCIAL", bloques=None):
        return TextoResolucionCertificado.objects.create(
            textoresolucion_nombre="Techo Digno - Nación - Parcial",
            textoresolucion_programa=self.programa,
            textoresolucion_financiamiento="N",
            textoresolucion_tipo=tipo,
            textoresolucion_bloques=bloques or [
                {"clase": "visto", "label": "", "texto": "El Expediente N° {{ certificado.expediente }};"},
                {"clase": "considerando", "label": "", "texto": "Que la Empresa {{ empresa.nombre }} certificó {{ certificado.monto_pesos|pesos }};"},
                {"clase": "resuelve", "label": "", "texto": ""},
                {"clase": "articulo", "label": "", "texto": "Apruébase por {{ certificado.monto_pesos|pesos }}."},
            ],
        )

    # -- permisos --------------------------------------------------------- #

    def test_anonimo_es_redirigido_al_login(self):
        for url in (reverse("carga:lista-textos-resolucion"),
                    reverse("carga:crear-texto-resolucion"),
                    reverse("carga:editar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})):
            assert self.client.get(url).status_code == 302

    def test_usuario_sin_permisos_recibe_403(self):
        self.client.force_login(self.sin_permisos)
        for url in (reverse("carga:lista-textos-resolucion"),
                    reverse("carga:crear-texto-resolucion"),
                    reverse("carga:editar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})):
            assert self.client.get(url).status_code == 403

    def test_superuser_accede(self):
        self._texto()
        self.client.force_login(self.superuser)
        assert self.client.get(reverse("carga:lista-textos-resolucion")).status_code == 200
        assert self.client.get(reverse("carga:crear-texto-resolucion")).status_code == 200

    # -- CRUD de la plantilla --------------------------------------------- #

    def _post_plantilla(self, textos, accion="guardar", extra=None):
        datos = {
            "textoresolucion_nombre": "Techo Digno - Nación",
            "textoresolucion_programa": self.programa.pk,
            "textoresolucion_financiamiento": "N",
            "textoresolucion_tipo": "",
            "bloques-TOTAL_FORMS": str(len(textos)),
            "bloques-INITIAL_FORMS": "0",
            "bloques-MIN_NUM_FORMS": "0",
            "bloques-MAX_NUM_FORMS": "1000",
            "accion": accion,
        }
        for i, (clase, texto) in enumerate(textos):
            datos[f"bloques-{i}-clase"] = clase
            datos[f"bloques-{i}-label"] = ""
            datos[f"bloques-{i}-texto"] = texto
        datos.update(extra or {})
        return self.client.post(reverse("carga:crear-texto-resolucion"), datos)

    def test_guardar_crea_la_plantilla_con_sus_bloques(self):
        self.client.force_login(self.superuser)
        response = self._post_plantilla([("visto", "El Expediente {{ certificado.expediente }};")])

        assert response.status_code == 302
        texto = TextoResolucionCertificado.objects.get()
        assert texto.textoresolucion_bloques == [
            {"clase": "visto", "label": "", "texto": "El Expediente {{ certificado.expediente }};"}
        ]

    def test_jinja_invalido_no_se_guarda(self):
        """Sin esta validación, un `{%` suelto rompería la ficha de todos los
        certificados del alcance."""
        self.client.force_login(self.superuser)
        response = self._post_plantilla([("considerando", "{% if x %}sin cerrar")])

        assert response.status_code == 200
        assert not TextoResolucionCertificado.objects.exists()
        assert "Error de sintaxis" in response.content.decode()

    def test_un_bloque_que_necesita_texto_no_puede_ir_vacio(self):
        self.client.force_login(self.superuser)
        response = self._post_plantilla([("considerando", "")])

        assert response.status_code == 200
        assert not TextoResolucionCertificado.objects.exists()

    def test_la_formula_resuelve_no_necesita_texto(self):
        self.client.force_login(self.superuser)
        response = self._post_plantilla([("resuelve", "")])

        assert response.status_code == 302
        assert TextoResolucionCertificado.objects.get().textoresolucion_bloques == [
            {"clase": "resuelve", "label": "", "texto": ""}
        ]

    def test_la_vista_previa_renderiza_contra_el_certificado_elegido(self):
        self.client.force_login(self.superuser)
        response = self._post_plantilla(
            [("considerando", "Que {{ empresa.nombre }} certificó {{ certificado.monto_pesos|pesos }};")],
            accion="previsualizar",
            extra={"certificado": str(self.certificado.pk)},
        )

        assert response.status_code == 200
        cuerpo = response.content.decode()
        assert "Constructora Sur SRL" in cuerpo
        assert "$1.500.000,00" in cuerpo
        # Previsualizar no guarda.
        assert not TextoResolucionCertificado.objects.exists()

    def test_la_vista_previa_avisa_que_faltan_datos(self):
        self.client.force_login(self.superuser)
        response = self._post_plantilla(
            [("considerando", "Que {{ inexistente.campo }};")],
            accion="previsualizar",
            extra={"certificado": str(self.certificado.pk)},
        )

        assert "Faltan datos" in response.content.decode()

    # -- flujo por certificado -------------------------------------------- #

    def test_sin_plantilla_ofrece_crearla_en_vez_de_404(self):
        self.client.force_login(self.superuser)
        response = self.client.get(
            reverse("carga:editar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})
        )

        assert response.status_code == 409
        assert "No hay ningún texto de resolución cargado" in response.content.decode()

    def test_muestra_el_texto_resuelto_desde_la_plantilla(self):
        self._texto()
        self.client.force_login(self.superuser)
        response = self.client.get(
            reverse("carga:editar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})
        )

        assert response.status_code == 200
        cuerpo = response.content.decode()
        assert "El Expediente N° EXP-CERT;" in cuerpo
        assert "$1.500.000,00" in cuerpo

    def test_guardar_deja_el_snapshot_con_su_procedencia(self):
        plantilla = self._texto()
        self.client.force_login(self.superuser)
        url = reverse("carga:editar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})

        response = self.client.post(url, {
            "bloques-TOTAL_FORMS": "1",
            "bloques-INITIAL_FORMS": "0",
            "bloques-MIN_NUM_FORMS": "0",
            "bloques-MAX_NUM_FORMS": "1000",
            "bloques-0-clase": "considerando",
            "bloques-0-label": "",
            "bloques-0-texto": "Texto retocado a mano;",
            "accion": "guardar",
        })

        assert response.status_code == 302
        snapshot = Certificado.objects.get(pk=self.certificado.pk).certificado_texto_resolucion
        assert snapshot["bloques"] == [{"clase": "considerando", "label": "", "texto": "Texto retocado a mano;"}]
        assert snapshot["textoresolucion_id"] == plantilla.pk
        assert snapshot["textoresolucion_history_id"] is not None

    def test_el_snapshot_gana_sobre_la_plantilla(self):
        self._texto()
        Certificado.objects.filter(pk=self.certificado.pk).update(
            certificado_texto_resolucion={"bloques": [{"clase": "considerando", "label": "", "texto": "Snapshot;"}]}
        )
        self.client.force_login(self.superuser)
        response = self.client.get(
            reverse("carga:editar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})
        )

        cuerpo = response.content.decode()
        assert "Snapshot;" in cuerpo
        assert "El Expediente N° EXP-CERT;" not in cuerpo

    def test_restaurar_borra_el_snapshot(self):
        self._texto()
        Certificado.objects.filter(pk=self.certificado.pk).update(
            certificado_texto_resolucion={"bloques": [{"clase": "considerando", "label": "", "texto": "Snapshot;"}]}
        )
        self.client.force_login(self.superuser)

        self.client.post(
            reverse("carga:restaurar-texto-resolucion-certificado", kwargs={"pk": self.certificado.pk})
        )

        assert Certificado.objects.get(pk=self.certificado.pk).certificado_texto_resolucion is None

    # -- generación del documento ----------------------------------------- #

    def _cargar_encabezado(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from secretariador.models import EncabezadoDocumento

        base = io.BytesIO()
        Document().save(base)
        base.seek(0)
        return EncabezadoDocumento.objects.create(
            encabezadodocumento_archivo=SimpleUploadedFile("encabezado.docx", base.read())
        )

    def test_genera_el_docx(self):
        self._texto()
        self._cargar_encabezado()
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse("carga:resolucion-certificado-docx", kwargs={"pk": self.certificado.pk})
        )

        assert response.status_code == 200
        assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert "EXP-CERT" in response["Content-Disposition"]
        textos = [p.text for p in Document(io.BytesIO(response.content)).paragraphs if p.text.strip()]
        assert "VISTO: El Expediente N° EXP-CERT;" in textos
        assert "Artículo 1º: Apruébase por $1.500.000,00." in textos

    def test_no_genera_el_docx_si_faltan_datos(self):
        """Una resolución se firma: emitirla con «falta: ...» adentro es peor que
        no emitirla."""
        self._texto(bloques=[{"clase": "considerando", "label": "", "texto": "Que {{ inexistente.campo }};"}])
        self._cargar_encabezado()
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse("carga:resolucion-certificado-docx", kwargs={"pk": self.certificado.pk})
        )

        assert response.status_code == 302
        assert "texto-resolucion" in response["Location"]

    def test_sin_encabezado_cargado_avisa_en_vez_de_reventar(self):
        self._texto()
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse("carga:resolucion-certificado-docx", kwargs={"pk": self.certificado.pk})
        )

        assert response.status_code == 500
        assert "documento base" in response.content.decode()


class ManualTextosResolucionTests(TestCase):
    def test_lista_el_catalogo_de_variables(self):
        """El manual se arma desde el mismo catálogo que la paleta del editor: si
        alguien agrega una variable, aparece acá sin tocar el HTML."""
        user = get_user_model().objects.create_user(username="lector", password="testpass123")
        self.client.force_login(user)

        response = self.client.get(reverse("carga:ayuda-textos-resolucion"))

        assert response.status_code == 200
        cuerpo = response.content.decode()
        assert "certificado.monto_pesos|pesos" in cuerpo
        assert "firmantes.gerente.nombre_completo" in cuerpo


class ContextoCertificadoVaciosTests(TestCase):
    """Un dato sin cargar tiene que reportarse como faltante, no renderizarse como
    cadena vacía: de lo contrario la resolución sale diciendo "del período ," y
    nadie se entera hasta que está firmada."""

    def test_descarta_las_claves_sin_valor(self):
        limpio = resolucion_texto._sin_vacios({
            "con_valor": "x", "nulo": None, "vacio": "", "cero": Decimal("0"),
            "anidado": {"nulo": None, "con_valor": "y"},
        })

        assert limpio == {"con_valor": "x", "cero": Decimal("0"), "anidado": {"con_valor": "y"}}

    def test_el_cero_no_se_descarta(self):
        # Un monto de 0 es un dato cargado, no un dato faltante.
        assert resolucion_texto._sin_vacios({"monto": Decimal("0")}) == {"monto": Decimal("0")}

    def test_un_campo_nulo_del_certificado_se_reporta_como_faltante(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-V"
        )
        rubro = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        CertificadoFinanciamiento.objects.create(
            certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N"
        )
        certificado = Certificado.objects.create(
            certificado_obra=obra, certificado_tipo="PARCIAL", certificado_financiamiento="N",
            certificado_rubro_db=rubro, certificado_expediente="EXP-C",
            certificado_periodo=None, certificado_fecha=date(2026, 3, 31),
        )

        contexto = resolucion_texto.contexto_certificado(certificado)
        texto, faltantes = resolucion_texto.render_bloque("del período {{ certificado.periodo }};", contexto)

        assert "«falta:" in texto
        assert faltantes == {"periodo"}


class HistorialPlanDeTrabajosTests(TestCase):
    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-H"
        )
        self.plan = PlanDeTrabajos.objects.create(trabajos_obra=obra, trabajos_fecha=date(2026, 1, 1))
        self.rubro = PlanDeTrabajosRubro.objects.create(
            rubro_plan=self.plan, rubro_nombre="Vivienda", rubro_presupuesto=Decimal("1000")
        )
        self.item = PlanDeTrabajosItem.objects.create(
            planitem_rubro=self.rubro, planitem_nombre="Mampostería", planitem_incidencia_pct=Decimal("100")
        )
        self.user = get_user_model().objects.create_superuser(
            username="editor", password="testpass123", first_name="Ana", last_name="Gómez"
        )
        self.url = reverse(
            "historial", kwargs={"app_label": "carga", "model_name": "plandetrabajos", "pk": self.plan.pk}
        )

    def test_muestra_el_diff_y_quien_hizo_el_cambio(self):
        self.rubro.rubro_presupuesto = Decimal("1500")
        self.rubro._history_user = self.user
        self.rubro.save()
        self.client.force_login(self.user)

        response = self.client.get(self.url)

        self.assertContains(response, "Ana Gómez")
        self.assertContains(response, "Presupuesto")
        self.assertContains(response, "1000.00")
        self.assertContains(response, "1500")

    def test_incluye_hijos_eliminados(self):
        self.item.delete()
        self.client.force_login(self.user)

        response = self.client.get(self.url)

        self.assertContains(response, "Eliminado")
        self.assertContains(response, "Mampostería · Vivienda")

    def test_omite_guardados_sin_cambios(self):
        from carga.history import plandetrabajos_sources
        from core.history import build_timeline

        self.rubro.save()
        timeline = build_timeline(plandetrabajos_sources(self.plan))
        tipos = [e["tipo"] for g in timeline["grupos"] for e in g["entradas"]]

        self.assertNotIn("Modificado", tipos)

    def test_la_pagina_del_plan_muestra_la_pestana_una_sola_vez(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("carga:update-plandetrabajos", kwargs={"pk": self.plan.pk}))

        self.assertContains(response, 'id="historialSidebar"', count=1)
        self.assertContains(response, self.url)

    def test_obra_con_excluded_fields_no_rompe_el_historial(self):
        obra = self.plan.trabajos_obra
        obra.obra_nombre = "Obra Renombrada"
        obra.save()
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("historial", kwargs={"app_label": "carga", "model_name": "obra", "pk": obra.pk})
        )

        self.assertContains(response, "Obra Renombrada")

    def test_la_matriz_de_etapas_muestra_el_historial_del_plan(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("carga:plandetrabajosetapa-matriz", kwargs={"pk": self.rubro.pk}))

        self.assertContains(response, self.url)

    def test_sin_permiso_da_403(self):
        sin_permiso = get_user_model().objects.create_user(username="nadie", password="testpass123")
        self.client.force_login(sin_permiso)

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 403)


class RetencionAdobeTests(TestCase):
    """Decreto 654/2015: tres por mil sobre el valor bruto de los certificados de obras que
    usan ladrillos de adobe como insumo. Se decide por el rubro del certificado (Vivienda),
    no por un dato de la Obra, para cubrir también los certificados legacy.

    A diferencia del Fondo de Reparo, el decreto no exceptúa a los certificados de
    Anticipo -- por eso hay un test explícito que lo confirma en vez de asumirlo."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.rubro_vivienda = CertificadoRubro.objects.create(
            certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V"
        )
        self.rubro_infraestructura = CertificadoRubro.objects.create(
            certificadorubro_nombre="Infraestructura Frentista", certificadorubro_nombre_corto="F"
        )
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-ADOBE",
        )

    def _certificado(self, rubro, tipo="PARCIAL", **kwargs):
        defaults = {
            "certificado_obra": self.obra,
            "certificado_tipo": tipo,
            "certificado_financiamiento": "N",
            "certificado_rubro_db": rubro,
            "certificado_expediente": "EXP",
            "certificado_fecha": date(2026, 6, 1),
            "certificado_monto_pesos": Decimal("100000"),
            "certificado_monto_uvi": Decimal("1000"),
        }
        defaults.update(kwargs)
        return Certificado(**defaults)

    def test_aplica_a_los_certificados_de_rubro_vivienda(self):
        certificado = self._certificado(self.rubro_vivienda)

        self.assertTrue(certificado.certificado_aplica_retencion_adobe)
        # tres por mil de 100.000 = 300; de 1.000 UVI = 3.
        self.assertEqual(certificado.certificado_retencion_adobe_monto_pesos(), Decimal("300"))
        self.assertEqual(certificado.certificado_retencion_adobe_monto_uvi(), Decimal("3"))

    def test_no_aplica_a_otros_rubros(self):
        certificado = self._certificado(self.rubro_infraestructura)

        self.assertFalse(certificado.certificado_aplica_retencion_adobe)
        self.assertEqual(certificado.certificado_retencion_adobe_monto_pesos(), Decimal("0"))
        self.assertEqual(certificado.certificado_retencion_adobe_monto_uvi(), Decimal("0"))

    def test_cubre_los_certificados_legacy(self):
        certificado = self._certificado(self.rubro_vivienda, tipo="LEGACY")

        self.assertEqual(certificado.certificado_retencion_adobe_monto_pesos(), Decimal("300"))

    def test_se_aplica_tambien_a_certificados_de_anticipo(self):
        """El decreto habla del "valor bruto de los certificados" sin exceptuar ningún
        tipo -- a propósito, esto NO replica la exclusión de Fondo de Reparo para ANTICIPO."""
        certificado = self._certificado(self.rubro_vivienda, tipo="ANTICIPO")

        self.assertEqual(certificado.certificado_retencion_adobe_monto_pesos(), Decimal("300"))

    def test_monto_certificado_nulo_no_rompe_el_calculo(self):
        certificado = self._certificado(self.rubro_vivienda, certificado_monto_pesos=None, certificado_monto_uvi=None)

        self.assertEqual(certificado.certificado_retencion_adobe_monto_pesos(), Decimal("0"))
        self.assertEqual(certificado.certificado_retencion_adobe_monto_uvi(), Decimal("0"))


class RetencionAdobeContextoTextoTests(TestCase):
    """La retención tiene que estar disponible como variable para armar el texto de la
    resolución (el pedido original: "tener como dato a la hora de generar los textos")."""

    def setUp(self):
        empresa = Empresa.objects.create(empresa_nombre="Empresa Test")
        programa = Programa.objects.create(programa_nombre="Programa Test")
        self.rubro_vivienda = CertificadoRubro.objects.create(certificadorubro_nombre="Vivienda", certificadorubro_nombre_corto="V")
        self.rubro_infraestructura = CertificadoRubro.objects.create(certificadorubro_nombre="Infraestructura Frentista", certificadorubro_nombre_corto="F")
        CertificadoFinanciamiento.objects.create(certificadofinanciamiento_nombre="Nación", certificadofinanciamiento_nombre_corto="N")
        self.obra = Obra.objects.create(
            obra_nombre="Obra Test", obra_empresa=empresa, obra_programa=programa, obra_expediente="EXP-CTX",
        )

    def _certificado(self, rubro):
        return Certificado.objects.create(
            certificado_obra=self.obra, certificado_tipo="PARCIAL", certificado_financiamiento="N",
            certificado_rubro_db=rubro, certificado_expediente="EXP-1", certificado_fecha=date(2026, 6, 1),
            certificado_monto_pesos=Decimal("100000"),
        )

    def test_el_contexto_trae_la_retencion(self):
        contexto = resolucion_texto.contexto_certificado(self._certificado(self.rubro_vivienda))

        self.assertIs(contexto["certificado"]["aplica_retencion_adobe"], True)
        self.assertEqual(contexto["certificado"]["retencion_adobe_pesos"], Decimal("300"))

    def test_se_puede_condicionar_un_bloque_con_la_variable_booleana(self):
        contexto = resolucion_texto.contexto_certificado(self._certificado(self.rubro_vivienda))
        texto, faltantes = resolucion_texto.render_bloque(
            "{% if certificado.aplica_retencion_adobe %}"
            "Que corresponde retener {{ certificado.retencion_adobe_pesos|pesos }} en concepto del "
            "Decreto 654/2015.{% endif %}",
            contexto,
        )

        self.assertEqual(texto, "Que corresponde retener $300,00 en concepto del Decreto 654/2015.")
        self.assertEqual(faltantes, set())

    def test_en_otro_rubro_el_cero_no_se_reporta_como_dato_faltante(self):
        """Fuera de Vivienda la retención es 0 -- un valor real, no un dato sin cargar --
        así que no debe aparecer en `faltantes` ni bloquear la generación del .docx."""
        contexto = resolucion_texto.contexto_certificado(self._certificado(self.rubro_infraestructura))
        texto, faltantes = resolucion_texto.render_bloque("{{ certificado.retencion_adobe_pesos|pesos }}", contexto)

        self.assertIs(contexto["certificado"]["aplica_retencion_adobe"], False)
        self.assertEqual(texto, "$0,00")
        self.assertEqual(faltantes, set())
